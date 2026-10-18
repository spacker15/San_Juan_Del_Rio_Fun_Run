"""Build the test workbooks, recalculate them in one LibreOffice call, share the results."""
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "build"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import build_workbook  # noqa: E402
import demo_data  # noqa: E402
import scenarios  # noqa: E402

SCENARIOS = {"blank": dict, "demo": demo_data.generate, "ties": scenarios.ties}


class Book:
    def __init__(self, path, data):
        self.wb = openpyxl.load_workbook(path, data_only=True)
        self.data = data

    def __getitem__(self, sheet):
        return self.wb[sheet]

    def find(self, sheet, col, text):
        ws = self.wb[sheet]
        for r in range(1, ws.max_row + 1):
            if str(ws[f"{col}{r}"].value or "") == text:
                return r
        raise KeyError(text)


@pytest.fixture(scope="session")
def books():
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice (soffice) is required to recalculate formulas")
    tmp = tempfile.mkdtemp(prefix="srw_")
    out = os.path.join(tmp, "out")
    datas, paths = {}, []
    for key, make in SCENARIOS.items():
        datas[key] = make()
        paths.append(build_workbook.build(os.path.join(tmp, f"{key}.xlsx"), datas[key]))
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, *paths],
                   check=True, capture_output=True, timeout=900)
    return {k: Book(os.path.join(out, f"{k}.xlsx"), datas[k]) for k in SCENARIOS}
