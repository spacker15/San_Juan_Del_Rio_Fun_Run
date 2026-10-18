#!/usr/bin/env python3
"""Generate the Saints Run Wild fundraising workbook (.xlsx, no macros).

Usage:
    python build/build_workbook.py            # blank template + demo workbook
    python build/build_workbook.py --blank    # template only

Sheet groups
    Outputs (red tabs):   TV Dashboard, Leaderboards, Challenge Winners, Prize Report, Data Check
    Inputs  (blue tabs):  Settings, Classrooms, Roster, Prizes, Challenges, OneCause_Paste, Cash_Check, History
    Reports (green tabs): Rpt_* print-ready sheets
    Engine  (hidden):     OC_Norm, Ledger, Calc_* (all math lives here; outputs only read it)

Formulas stick to Excel 2010+ functions (SUMIFS/COUNTIFS/INDEX/MATCH/LARGE/SMALL and
CSE array formulas) so the file works in Excel desktop, Excel for the web, Mac, and LibreOffice.
"""
import argparse
import os
import sys
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config as C  # noqa: E402
import styles as S  # noqa: E402

# ----------------------------------------------------------------------------
# Geometry
# ----------------------------------------------------------------------------
STU_LAST = 1 + C.N_STU
CLS_LAST = 1 + C.N_CLS
OC_LAST = 1 + C.N_OC
CC_LAST = 1 + C.N_CC
LG_LAST = 1 + C.N_OC + C.N_CC
CH_LAST = 1 + C.N_CH
PR_LAST = 1 + C.N_PR
G_LAST = 1 + len(C.GRADES)
DAY_LAST = 1 + C.N_DAY
FAM_LAST = 1 + C.N_FAM
DON_LAST = 1 + C.N_DON

# Calc_Students per-challenge columns
WIN = [CL(35 + i) for i in range(C.N_CH)]    # $ inside challenge window
ASOF = [CL(50 + i) for i in range(C.N_CH)]   # cumulative $ as of challenge end
# Calc_Classes / Calc_Grades per-challenge metric columns
CMET = [CL(20 + i) for i in range(C.N_CH)]
GMET = [CL(14 + i) for i in range(C.N_CH)]
# Calc_Challenges winner blocks
GB_ROW0, CB_ROW0 = 20, 40                     # grade block rows 20-30, class block rows 40-79
BW = [CL(2 + 2 * i) for i in range(C.N_CH)]  # winner text column per challenge
BA = [CL(3 + 2 * i) for i in range(C.N_CH)]  # amount column per challenge
GB_LAST = GB_ROW0 + len(C.GRADES) - 1
CB_LAST = CB_ROW0 + C.N_CLS - 1


def rng(sheet, col, r1, r2):
    return f"'{sheet}'!${col}${r1}:${col}${r2}"


def cs(col):
    return rng("Calc_Students", col, 2, STU_LAST)


def ccl(col):
    return rng("Calc_Classes", col, 2, CLS_LAST)


def cg(col):
    return rng("Calc_Grades", col, 2, G_LAST)


def lg(col):
    return rng("Ledger", col, 2, LG_LAST)


def oc(col):
    return rng("OC_Norm", col, 2, OC_LAST)


def ch(col):
    return rng("Calc_Challenges", col, 2, CH_LAST)


def cc(col):
    return rng("Cash_Check", col, 2, CC_LAST)


# ----------------------------------------------------------------------------
# Small styling helpers
# ----------------------------------------------------------------------------
def put(ws, cell, value=None, *, font=None, fill=None, fmt=None, align=None, border=None):
    c = ws[cell]
    if value is not None:
        c.value = value
    if font:
        c.font = font
    if fill:
        c.fill = S.fill(fill) if isinstance(fill, str) else fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def arr(ws, cell, formula):
    ws[cell] = ArrayFormula(cell, formula)


def header_row(ws, row, labels, col0=1, fill=S.NAVY, height=30):
    for i, lab in enumerate(labels):
        put(ws, f"{CL(col0 + i)}{row}", lab, font=S.HEADER_FONT, fill=fill, align=S.CENTER, border=S.BOX)
    ws.row_dimensions[row].height = height


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def banner(ws, cells, text, size=20, fill=S.NAVY, color=S.WHITE, height=34):
    ws.merge_cells(cells)
    first = cells.split(":")[0]
    put(ws, first, text, font=S.display(size, color), fill=fill, align=S.CENTER)
    row = int("".join(ch for ch in first if ch.isdigit()))
    ws.row_dimensions[row].height = height


def iter_cells(ws, cells):
    if ":" not in cells:
        cells = f"{cells}:{cells}"
    return ws[cells]


def style_range(ws, cells, *, fill=None, fmt=None, font=None, border=None, align=None):
    for row in iter_cells(ws, cells):
        for c in row:
            if fill:
                c.fill = S.fill(fill)
            if fmt:
                c.number_format = fmt
            if font:
                c.font = font
            if border:
                c.border = border
            if align:
                c.alignment = align


def row_lines(ws, cells, key_cell):
    """Light row rules only where the row has data, so blank formula rows print clean."""
    ws.conditional_formatting.add(cells, FormulaRule(formula=[f'{key_cell}<>""'], border=S.BOX))


def unlock(ws, cells):
    for row in iter_cells(ws, cells):
        for c in row:
            c.protection = Protection(locked=False)


def protect(ws):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def add_list_dv(ws, cells, formula, prompt=None, strict=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showErrorMessage=True, errorStyle="stop" if strict else "warning")
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "Tip", prompt, True
    ws.add_data_validation(dv)
    dv.add(cells)


def name(wb, nm, target):
    wb.defined_names[nm] = DefinedName(nm, attr_text=target)


def print_setup(ws, landscape=False, fit_width=True, title_rows=None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    if fit_width:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.left.text = "&A"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddHeader.right.text = "Printed &D"
    if title_rows:
        ws.print_title_rows = title_rows


def dynamic_print_area(ws, last_col, count_formula, header_rows):
    """Print area that grows with the data, so blank formula rows don't print as empty pages."""
    t = f"'{ws.title}'"
    ws.defined_names["_xlnm.Print_Area"] = DefinedName(
        "_xlnm.Print_Area",
        attr_text=f"{t}!$A$1:INDEX({t}!${last_col}:${last_col},{header_rows}+MAX(1,{count_formula}))",
    )


# ----------------------------------------------------------------------------
# INPUT SHEETS
# ----------------------------------------------------------------------------
def build_settings(wb, ws, settings):
    ws.sheet_properties.tabColor = S.TAB_INPUT
    widths(ws, {"A": 34, "B": 38, "C": 16, "D": 22, "E": 3, "H": 12, "I": 32, "J": 28, "K": 10, "L": 8})
    banner(ws, "A1:D1", "SETTINGS  •  change these first", 16)
    rows = [
        (3, "Campaign name", settings["campaign"], None),
        (4, "School", settings["school"], None),
        (5, "Campaign year", settings["year"], "0"),
        (6, "Fundraising goal ($)", settings["goal"], S.FMT_USD),
        (7, "Stretch milestone ($)", settings["stretch"], S.FMT_USD),
        (8, "Stretch milestone label", settings["stretch_label"], None),
        (9, "Campaign start date", settings["start"], S.FMT_DATE),
        (10, "Campaign end date (Fun Run)", settings["end"], S.FMT_DATE),
        (11, "Participation minimum ($ per student)", settings["part_min"], S.FMT_USD2),
        (12, "Public name display (TV / Leaderboards)", settings["privacy"], None),
        (13, "TV tagline", settings["tagline"], None),
        (14, "Compare to year (History tab)", settings["comp_year"], "0"),
        (15, "'As of' date override (blank = today)", settings["asof"], S.FMT_DATE),
    ]
    for r, label, val, fmt in rows:
        put(ws, f"A{r}", label, font=S.font(11, True), border=S.BOX)
        put(ws, f"B{r}", val, fill=S.INPUT, border=S.BOX, fmt=fmt, align=S.LEFT)
    put(ws, "C15", "Leave blank in real use", font=S.font(9, italic=True, color="777777"))
    add_list_dv(ws, "B12", "=PrivacyList")

    put(ws, "A17", "ONECAUSE EXPORT MAPPING", font=S.display(12, S.RED))
    put(ws, "B17", "Type the exact column header from the OneCause export", font=S.font(9, italic=True))
    header_row(ws, 18, ["Field", "OneCause header text", "Column # found", "Status"], 1)
    for i, (field, hdr, required) in enumerate(C.ONECAUSE_MAP):
        r = 19 + i
        put(ws, f"A{r}", field + (" *" if required else ""), border=S.BOX)
        put(ws, f"B{r}", hdr or None, fill=S.INPUT, border=S.BOX)
        put(ws, f"C{r}", f"=IF(TRIM(B{r})=\"\",0,IFERROR(MATCH(TRIM(B{r}),'OneCause_Paste'!$1:$1,0),0))",
            fill=S.CALC, border=S.BOX, align=S.CENTER)
        if required:
            st = f"=IF(C{r}>0,\"✔ found\",\"✖ NOT FOUND, fix header\")"
        else:
            st = f"=IF(TRIM(B{r})=\"\",\"(not used)\",IF(C{r}>0,\"✔ found\",\"✖ not found\"))"
        put(ws, f"D{r}", st, fill=S.CALC, border=S.BOX)
    put(ws, "A28", "* required. Student name needs either the full-name header OR both first/last headers.",
        font=S.font(9, italic=True))

    # Dropdown source lists
    put(ws, "H2", "Grades", font=S.font(10, True))
    for i, g in enumerate(C.GRADES):
        ws[f"H{3 + i}"] = g
        ws[f"H{3 + i}"].number_format = "@"
    put(ws, "I2", "Challenge metric types", font=S.font(10, True))
    for i, m in enumerate(C.METRICS):
        ws[f"I{3 + i}"] = m
    put(ws, "J2", "Privacy options", font=S.font(10, True))
    for i, p in enumerate(C.PRIVACY_OPTIONS):
        ws[f"J{3 + i}"] = p
    put(ws, "K2", "Methods", font=S.font(10, True))
    for i, m in enumerate(C.METHODS):
        ws[f"K{3 + i}"] = m
    put(ws, "L2", "Y/N", font=S.font(10, True))
    ws["L3"], ws["L4"] = "Y", "N"
    for col in "HIJKL":
        ws.column_dimensions[col].hidden = True

    for nm, cell in [("CampaignName", "B3"), ("SchoolName", "B4"), ("CampaignYear", "B5"), ("Goal", "B6"),
                     ("Stretch", "B7"), ("StretchLabel", "B8"), ("StartDate", "B9"), ("EndDate", "B10"),
                     ("PartMin", "B11"), ("Privacy", "B12"), ("Tagline", "B13"), ("CompYear", "B14")]:
        name(wb, nm, f"'Settings'!${cell[0]}${cell[1:]}")
    name(wb, "GradeList", f"'Settings'!$H$3:$H${2 + len(C.GRADES)}")
    name(wb, "MetricList", f"'Settings'!$I$3:$I${2 + len(C.METRICS)}")
    name(wb, "PrivacyList", f"'Settings'!$J$3:$J${2 + len(C.PRIVACY_OPTIONS)}")
    name(wb, "MethodList", f"'Settings'!$K$3:$K${2 + len(C.METHODS)}")
    name(wb, "YNList", "'Settings'!$L$3:$L$4")
    keys = ["OC_ColDate", "OC_ColAmt", "OC_ColName", "OC_ColFirst", "OC_ColLast", "OC_ColDonor", "OC_ColTxn",
            "OC_ColPay"]
    for i, k in enumerate(keys):
        name(wb, k, f"'Settings'!$C${19 + i}")
    ws.freeze_panes = "A3"


def build_classrooms(ws, classrooms):
    ws.sheet_properties.tabColor = S.TAB_INPUT
    header_row(ws, 1, ["Teacher (as students know them)", "Grade", "Classroom (auto)", "Class goal $ (optional)",
                       "Notes"])
    widths(ws, {"A": 30, "B": 9, "C": 32, "D": 18, "E": 30})
    for r in range(2, CLS_LAST + 1):
        for col in "ABDE":
            put(ws, f"{col}{r}", fill=S.INPUT, border=S.BOX)
        ws[f"B{r}"].number_format = "@"
        ws[f"D{r}"].number_format = S.FMT_USD
        put(ws, f"C{r}", f"=IF(TRIM(A{r})=\"\",\"\",TRIM(A{r})&\" – \"&B{r})", fill=S.CALC, border=S.BOX)
    for i, (teacher, grade, goal) in enumerate(classrooms):
        r = 2 + i
        ws[f"A{r}"], ws[f"B{r}"] = teacher, str(grade)
        if goal:
            ws[f"D{r}"] = goal
    add_list_dv(ws, f"B2:B{CLS_LAST}", "=GradeList")
    ws.freeze_panes = "A2"


def build_roster(ws, roster):
    ws.sheet_properties.tabColor = S.TAB_INPUT
    header_row(ws, 1, ["Student ID (fixed)", "First name", "Last name", "Grade", "Teacher", "Family / household",
                       "Parent email", "OneCause name (if different)", "Active (Y/N)", "Classroom (auto)",
                       "Check (auto)"])
    widths(ws, {"A": 11, "B": 16, "C": 18, "D": 8, "E": 22, "F": 22, "G": 26, "H": 24, "I": 9, "J": 28, "K": 30})
    for r in range(2, STU_LAST + 1):
        put(ws, f"A{r}", f"S{r - 1:04d}", fill=S.CALC, border=S.BOX)
        for col in "BCDEFGHI":
            put(ws, f"{col}{r}", fill=S.INPUT, border=S.BOX)
        ws[f"D{r}"].number_format = "@"
        put(ws, f"J{r}", f"=IF(TRIM(E{r})=\"\",\"\",TRIM(E{r})&\" – \"&D{r})", fill=S.CALC, border=S.BOX)
        put(ws, f"K{r}", (f"=IF(TRIM(B{r}&C{r})=\"\",\"\",IF(OR(D{r}&\"\"=\"\",TRIM(E{r})=\"\"),"
                          f"\"Needs grade + teacher\",IF(ISNA(MATCH(J{r},{ccl('B')},0)),"
                          f"\"Teacher/grade not on Classrooms\",\"✔\")))"), fill=S.CALC, border=S.BOX)
    for i, s in enumerate(roster):
        r = 2 + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = s["first"], s["last"], str(s["grade"])
        ws[f"E{r}"], ws[f"F{r}"] = s["teacher"], s.get("family")
        ws[f"G{r}"], ws[f"H{r}"], ws[f"I{r}"] = s.get("email"), s.get("alias"), s.get("active")
    add_list_dv(ws, f"D2:D{STU_LAST}", "=GradeList")
    add_list_dv(ws, f"E2:E{STU_LAST}", f"='Classrooms'!$A$2:$A${CLS_LAST}", strict=False)
    add_list_dv(ws, f"I2:I{STU_LAST}", "=YNList")
    ws.conditional_formatting.add(f"K2:K{STU_LAST}", FormulaRule(
        formula=[f'AND(K2<>"",K2<>"✔")'], fill=S.fill("F8D7DA"), font=S.font(11, True, S.RED)))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:K{STU_LAST}"


def build_prizes(ws, prizes):
    ws.sheet_properties.tabColor = S.TAB_INPUT
    header_row(ws, 1, ["Prize name", "Raise at least ($)", "Description", "Students earned (auto)"])
    widths(ws, {"A": 30, "B": 18, "C": 40, "D": 20})
    for r in range(2, PR_LAST + 1):
        for col in "ABC":
            put(ws, f"{col}{r}", fill=S.INPUT, border=S.BOX)
        ws[f"B{r}"].number_format = S.FMT_USD
        put(ws, f"D{r}", f"=IF(N(B{r})<=0,\"\",COUNTIFS({cs('A')},1,{cs('T')},\">=\"&{r - 1}))",
            fill=S.CALC, border=S.BOX, align=S.CENTER)
    for i, (nm, thr, desc) in enumerate(prizes):
        ws[f"A{2 + i}"], ws[f"B{2 + i}"], ws[f"C{2 + i}"] = nm, thr, desc
    put(ws, f"A{PR_LAST + 2}", "List prizes from smallest to largest threshold. Students earn every prize up "
        "to their total.", font=S.font(9, italic=True))


def build_challenges(ws, challenges):
    ws.sheet_properties.tabColor = S.TAB_INPUT
    header_row(ws, 1, ["Start date", "End date", "Day", "Challenge", "Description", "Reward", "Metric type",
                       "Min $ (optional)", "Target (optional)"], height=32)
    widths(ws, {"A": 12, "B": 12, "C": 12, "D": 30, "E": 46, "F": 40, "G": 30, "H": 11, "I": 12})
    for r in range(2, CH_LAST + 1):
        for col in "ABCDEFGHI":
            put(ws, f"{col}{r}", fill=S.INPUT, border=S.BOX)
        ws[f"A{r}"].number_format = ws[f"B{r}"].number_format = S.FMT_DATE_S
    for i, row in enumerate(challenges):
        r = 2 + i
        for col, v in zip("ABCDEFGHI", row):
            ws[f"{col}{r}"] = v
    add_list_dv(ws, f"G2:G{CH_LAST}", "=MetricList")
    notes = [
        "Metric types:",
        "CLASS_TOTAL_PER_GRADE: top class $ in each grade during the dates",
        "CLASS_PARTICIPATION_MIN: class with the highest % of students raising at least Min $ during the dates",
        "CLASS_PARTICIPATION_CUM: class with the highest overall participation % at the end of the date",
        "DRAWING: random pick of 'Target' students who got any gift during the dates (Challenge Winners tab)",
        "TOP_STUDENT_PER_CLASS: top student $ in each class during the dates",
        "GRADE_TOTAL: grade with the most $ during the dates",
        "SCHOOL_PARTICIPATION_TARGET: HIT when school participation reaches Target (e.g. 0.5 = 50%)",
        "SCHOOL_TOTAL_TARGET: HIT when school $ reaches Target by the end date",
        "NONE: display only",
    ]
    for i, t in enumerate(notes):
        put(ws, f"D{CH_LAST + 2 + i}", t, font=S.font(9, i == 0, italic=i > 0))
    ws.freeze_panes = "A2"


def build_onecause_paste(ws, headers, rows):
    ws.sheet_properties.tabColor = S.TAB_INPUT
    for i, h in enumerate(headers):
        put(ws, f"{CL(i + 1)}1", h, font=S.font(11, True))
    for r, row in enumerate(rows, start=2):
        for i, v in enumerate(row):
            if v is not None and v != "":
                ws.cell(r, i + 1, v)
    for i in range(1, 15):
        ws.column_dimensions[CL(i)].width = 18
    ws.freeze_panes = "A2"


def build_cash(ws, cash):
    ws.sheet_properties.tabColor = S.TAB_INPUT
    header_row(ws, 1, ["Date received", "Student (pick from list)", "Donor / family", "Amount", "Method",
                       "Check #", "Notes", "Entered by", "Student ID (auto)", "Status (auto)"], height=32)
    widths(ws, {"A": 13, "B": 44, "C": 24, "D": 11, "E": 10, "F": 10, "G": 24, "H": 12, "I": 12, "J": 14,
                "K": 4, "L": 4})
    for r in range(2, CC_LAST + 1):
        for col in "ABCDEFGH":
            put(ws, f"{col}{r}", fill=S.INPUT, border=S.BOX)
        ws[f"A{r}"].number_format = "m/d/yyyy"
        ws[f"D{r}"].number_format = S.FMT_USD2
        ws[f"I{r}"] = (f"=IF(TRIM(B{r})=\"\",\"\",IF(ISNUMBER(FIND(\"#\",B{r})),"
                       f"TRIM(MID(B{r},FIND(\"#\",B{r})+1,20)),UPPER(TRIM(B{r}))))")
        ws[f"J{r}"] = (f"=IF(AND(A{r}=\"\",B{r}=\"\",D{r}=\"\"),\"\",IF(OR(NOT(ISNUMBER(D{r})),N(D{r})<=0),"
                       f"\"BAD AMOUNT\",IF(NOT(ISNUMBER(A{r})),\"BAD DATE\",IF(I{r}=\"\",\"NO STUDENT\","
                       f"IF(ISNA(MATCH(I{r},{cs('B')},0)),\"NO STUDENT\",\"OK\")))))")
        ws[f"K{r}"] = f"=IF(J{r}=\"OK\",ROUND(D{r},2),0)"
        ws[f"L{r}"] = f"=IF(J{r}=\"OK\",INT(A{r}),\"\")"
        for col in "IJ":
            ws[f"{col}{r}"].fill = S.fill(S.CALC)
            ws[f"{col}{r}"].border = S.BOX
    ws.column_dimensions["K"].hidden = True
    ws.column_dimensions["L"].hidden = True
    for i, g in enumerate(cash):
        r = 2 + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = g["date"], g["student"], g.get("donor"), g["amount"]
        ws[f"E{r}"], ws[f"F{r}"], ws[f"G{r}"], ws[f"H{r}"] = g.get("method"), g.get("check"), g.get("notes"), g.get(
            "by")
    add_list_dv(ws, f"B2:B{CC_LAST}", f"={cs('X')}", "Start typing is not supported in old Excel: pick from the "
                "list, or type the Student ID (e.g. S0042)", strict=False)
    add_list_dv(ws, f"E2:E{CC_LAST}", "=MethodList")
    ws.conditional_formatting.add(f"J2:J{CC_LAST}", FormulaRule(
        formula=['AND(J2<>"",J2<>"OK")'], fill=S.fill("F8D7DA"), font=S.font(11, True, S.RED)))
    ws.freeze_panes = "A2"


def build_history(ws, history):
    ws.sheet_properties.tabColor = S.TAB_INPUT
    banner(ws, "A1:P1", "HISTORY  •  past campaigns for year-over-year comparison", 14)
    header_row(ws, 3, ["Year", "Goal", "Raised", "Participation %", "Students"] + [f"Grade {g}" for g in C.GRADES])
    widths(ws, {CL(i): 11 for i in range(1, 17)})
    for r in range(4, 14):
        for i in range(1, 17):
            c = put(ws, f"{CL(i)}{r}", fill=S.INPUT, border=S.BOX)
            c.number_format = "0" if i in (1, 5) else (S.FMT_PCT if i == 4 else S.FMT_USD)
    put(ws, "A15", "Daily cumulative $ by campaign day (day 1 = first day of that year's campaign)",
        font=S.font(10, True))
    put(ws, "A16", "Day #", font=S.HEADER_FONT, fill=S.NAVY, align=S.CENTER)
    for i in range(2, 7):
        put(ws, f"{CL(i)}16", fill=S.INPUT, border=S.BOX, fmt="0", font=S.font(11, True), align=S.CENTER)
    for d in range(1, C.N_DAY + 1):
        put(ws, f"A{16 + d}", d, align=S.CENTER, border=S.BOX)
        for i in range(2, 7):
            put(ws, f"{CL(i)}{16 + d}", fill=S.INPUT, border=S.BOX, fmt=S.FMT_USD)
    for i, row in enumerate(history.get("summary", [])):
        for j, v in enumerate(row):
            ws.cell(4 + i, 1 + j, v)
    for i, (year, series) in enumerate(history.get("daily", {}).items()):
        ws.cell(16, 2 + i, year)
        for d, v in enumerate(series):
            ws.cell(17 + d, 2 + i, v)
    put(ws, "H15", "Fill Goal/Raised/grade totals in row 4+. Year cells across row 16 label the daily columns.",
        font=S.font(9, italic=True))
    ws.freeze_panes = "B4"


# ----------------------------------------------------------------------------
# ENGINE SHEETS (hidden)
# ----------------------------------------------------------------------------
def build_oc_norm(ws):
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    header_row(ws, 1, ["Raw date", "Raw amount", "Student name", "Donor", "Txn ID", "Payment", "Date", "Amount",
                       "Name key", "Roster row", "Student ID", "Status", "Counted $", "Counted date"])

    def pull(colname, r):
        idx = f"INDEX('OneCause_Paste'!$A${r}:$AZ${r},1,{colname})"
        return f"IF({colname}=0,\"\",IF({idx}=\"\",\"\",{idx}))"

    for r in range(2, OC_LAST + 1):
        ws[f"A{r}"] = "=" + pull("OC_ColDate", r)
        ws[f"B{r}"] = "=" + pull("OC_ColAmt", r)
        ws[f"C{r}"] = (f"=IF(OC_ColName>0,{pull('OC_ColName', r)},IF(AND(OC_ColFirst>0,OC_ColLast>0),"
                       f"TRIM({pull('OC_ColFirst', r)}&\" \"&{pull('OC_ColLast', r)}),\"\"))")
        ws[f"D{r}"] = "=" + pull("OC_ColDonor", r)
        ws[f"E{r}"] = "=" + pull("OC_ColTxn", r)
        ws[f"F{r}"] = "=" + pull("OC_ColPay", r)
        ws[f"G{r}"] = (f"=IF(A{r}=\"\",\"\",IF(ISNUMBER(A{r}),INT(A{r}),IFERROR(INT(DATEVALUE(A{r})),"
                       f"IFERROR(INT(VALUE(A{r})),\"BAD\"))))")
        ws[f"H{r}"] = (f"=IF(B{r}=\"\",\"\",IF(ISNUMBER(B{r}),ROUND(B{r},2),IFERROR(ROUND(VALUE(SUBSTITUTE("
                       f"SUBSTITUTE(SUBSTITUTE(B{r},\"$\",\"\"),\",\",\"\"),\" \",\"\")),2),\"BAD\")))")
        ws[f"I{r}"] = f"=LOWER(TRIM(C{r}))"
        ws[f"J{r}"] = (f"=IF(I{r}=\"\",\"\",IFERROR(MATCH(I{r},{cs('K')},0),IFERROR(MATCH(I{r},{cs('I')},0),"
                       f"IFERROR(MATCH(I{r},{cs('J')},0),\"\"))))")
        ws[f"K{r}"] = f"=IF(J{r}=\"\",\"\",INDEX({cs('B')},J{r}))"
        ws[f"L{r}"] = (f"=IF(AND(A{r}=\"\",B{r}=\"\",C{r}=\"\"),\"\",IF(NOT(ISNUMBER(H{r})),\"BAD AMOUNT\","
                       f"IF(H{r}<0,\"NEGATIVE / REFUND\",IF(H{r}=0,\"ZERO AMOUNT\",IF(NOT(ISNUMBER(G{r})),"
                       f"\"BAD DATE\",IF(K{r}=\"\",\"UNMATCHED\",IF(E{r}=\"\",\"OK\",IF(MATCH(E{r},"
                       f"$E$2:$E${OC_LAST},0)<ROW()-1,\"DUPLICATE TXN\",\"OK\"))))))))")
        ws[f"M{r}"] = f"=IF(L{r}=\"OK\",H{r},0)"
        ws[f"N{r}"] = f"=IF(L{r}=\"OK\",G{r},\"\")"


def build_ledger(ws):
    """Single stacked ledger: OneCause rows then Cash_Check rows. Every total reads from here."""
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    header_row(ws, 1, ["Source", "Date", "Student ID", "Donor", "Counted $", "Status", "Donor key",
                       "First donor row", "Donor #"])
    for r in range(2, LG_LAST + 1):
        if r <= OC_LAST:
            ws[f"A{r}"] = "Online"
            ws[f"B{r}"] = f"='OC_Norm'!N{r}"
            ws[f"C{r}"] = f"='OC_Norm'!K{r}"
            ws[f"D{r}"] = f"=TRIM('OC_Norm'!D{r})"
            ws[f"E{r}"] = f"='OC_Norm'!M{r}"
            ws[f"F{r}"] = f"='OC_Norm'!L{r}"
        else:
            c = r - C.N_OC
            ws[f"A{r}"] = f"=IF(TRIM('Cash_Check'!E{c})=\"\",\"Offline\",TRIM('Cash_Check'!E{c}))"
            ws[f"B{r}"] = f"='Cash_Check'!L{c}"
            ws[f"C{r}"] = f"='Cash_Check'!I{c}"
            ws[f"D{r}"] = f"=TRIM('Cash_Check'!C{c})"
            ws[f"E{r}"] = f"='Cash_Check'!K{c}"
            ws[f"F{r}"] = f"='Cash_Check'!J{c}"
        ws[f"G{r}"] = f"=IF(AND(F{r}=\"OK\",D{r}<>\"\"),LOWER(D{r}),\"\")"
        ws[f"H{r}"] = f"=IF(G{r}=\"\",0,IF(MATCH(G{r},$G$2:$G${LG_LAST},0)=ROW()-1,1,0))"
        ws[f"I{r}"] = f"=H{r}" if r == 2 else f"=I{r - 1}+H{r}"


def build_calc_students(ws):
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    labels = ["Valid", "ID", "First", "Last", "Grade", "Teacher", "Classroom", "Family", "Key first last",
              "Key last, first", "Key alias", "Online $", "Offline $", "Total $", "Gifts", "Participating",
              "Sort key", "Rank", "Display name", "Prize level", "Highest prize", "Next prize $", "$ to next",
              "Pick label", "Grade order", "Class order", "First gift date", "Leaderboard filter key",
              "Prize report key", "Report order key", "First family row", "Family #", "Full name", "Draw key"]
    labels += [f"Win $ ch{i + 1}" for i in range(C.N_CH)] + [f"As-of $ ch{i + 1}" for i in range(C.N_CH)]
    header_row(ws, 1, labels)
    for r in range(2, STU_LAST + 1):
        f = {}
        f["A"] = (f"=IF(AND(TRIM('Roster'!B{r}&'Roster'!C{r})<>\"\","
                  f"UPPER(LEFT(TRIM('Roster'!I{r})&\"Y\",1))<>\"N\"),1,0)")
        f["B"] = f"=IF(A{r}=1,'Roster'!A{r},\"\")"
        f["C"] = f"=IF(A{r}=1,TRIM('Roster'!B{r}),\"\")"
        f["D"] = f"=IF(A{r}=1,TRIM('Roster'!C{r}),\"\")"
        f["E"] = f"=IF(A{r}=1,TRIM('Roster'!D{r}&\"\"),\"\")"
        f["F"] = f"=IF(A{r}=1,TRIM('Roster'!E{r}),\"\")"
        f["G"] = f"=IF(A{r}=1,IF(F{r}=\"\",\"(No teacher)\",F{r}&\" – \"&E{r}),\"\")"
        f["H"] = f"=IF(A{r}=1,TRIM('Roster'!F{r}),\"\")"
        f["I"] = f"=IF(A{r}=1,LOWER(TRIM(C{r}&\" \"&D{r})),\"#\")"
        f["J"] = f"=IF(A{r}=1,LOWER(TRIM(D{r}&\", \"&C{r})),\"#\")"
        f["K"] = f"=IF(AND(A{r}=1,TRIM('Roster'!H{r})<>\"\"),LOWER(TRIM('Roster'!H{r})),\"#\")"
        f["L"] = f"=IF(A{r}=1,SUMIFS({lg('E')},{lg('C')},B{r},{lg('A')},\"Online\"),0)"
        f["M"] = f"=IF(A{r}=1,SUMIFS({lg('E')},{lg('C')},B{r},{lg('A')},\"<>Online\"),0)"
        f["N"] = f"=ROUND(L{r}+M{r},2)"
        f["O"] = f"=IF(A{r}=1,COUNTIFS({lg('C')},B{r},{lg('F')},\"OK\"),0)"
        f["P"] = f"=IF(AND(A{r}=1,N{r}>0,N{r}>=PartMin),1,0)"
        f["Q"] = f"=IF(A{r}=1,N{r}+(1000-ROW())/10^7,-1)"
        f["R"] = f"=IF(A{r}=1,COUNTIFS($A$2:$A${STU_LAST},1,$N$2:$N${STU_LAST},\">\"&N{r})+1,\"\")"
        f["S"] = (f"=IF(A{r}<>1,\"\",IF(Privacy=\"Full name\",C{r}&\" \"&D{r},IF(Privacy=\"Initials only\","
                  f"LEFT(C{r},1)&\". \"&LEFT(D{r},1)&\".\",C{r}&\" \"&LEFT(D{r},1)&\".\")))")
        f["T"] = (f"=IF(A{r}=1,COUNTIFS('Prizes'!$B$2:$B${PR_LAST},\"<=\"&N{r},"
                  f"'Prizes'!$B$2:$B${PR_LAST},\">0\"),0)")
        f["U"] = f"=IF(T{r}>0,INDEX('Prizes'!$A$2:$A${PR_LAST},T{r}),\"\")"
        f["V"] = f"=IF(AND(A{r}=1,T{r}<PrizeCount),INDEX('Prizes'!$B$2:$B${PR_LAST},T{r}+1),\"\")"
        f["W"] = f"=IF(V{r}=\"\",\"\",V{r}-N{r})"
        f["X"] = f"=IF(A{r}=1,D{r}&\", \"&C{r}&\" — \"&G{r}&\" #\"&B{r},\"\")"
        f["Y"] = f"=IF(A{r}=1,IFERROR(MATCH(E{r},GradeList,0),99),99)"
        f["Z"] = f"=IF(A{r}=1,IFERROR(MATCH(G{r},{ccl('B')},0),99),99)"
        f["AB"] = f"=IF(AND(A{r}=1,OR(LB_Grade&\"\"=\"All grades\",E{r}=LB_Grade&\"\")),Q{r},-1)"
        f["AC"] = f"=IF(AND(A{r}=1,PR_SelLevel>0,T{r}>=PR_SelLevel),Y{r}*1000000+Z{r}*10000+ROW(),\"\")"
        f["AD"] = f"=IF(A{r}=1,Y{r}*1000000+Z{r}*10000+ROW(),\"\")"
        f["AE"] = f"=IF(AND(A{r}=1,H{r}<>\"\"),IF(MATCH(H{r},$H$2:$H${STU_LAST},0)=ROW()-1,1,0),0)"
        f["AF"] = f"=AE{r}" if r == 2 else f"=AF{r - 1}+AE{r}"
        f["AG"] = f"=IF(A{r}=1,C{r}&\" \"&D{r},\"\")"
        f["AH"] = (f"=IF(AND(A{r}=1,CH_SelType=\"DRAWING\"),IF(INDEX(${WIN[0]}{r}:${WIN[-1]}{r},1,CH_Sel)>0,"
                   f"RAND(),-1),-1)")
        for i in range(C.N_CH):
            k = i + 2  # Calc_Challenges row
            chs = "'Calc_Challenges'!"
            f[WIN[i]] = (f"=IF(AND(A{r}=1,{chs}$I${k}=1),ROUND(SUMIFS({lg('E')},{lg('C')},B{r},{lg('B')},"
                         f"\">=\"&{chs}$B${k},{lg('B')},\"<=\"&{chs}$C${k}),2),0)")
            f[ASOF[i]] = (f"=IF(AND(A{r}=1,{chs}$J${k}=1),ROUND(SUMIFS({lg('E')},{lg('C')},B{r},{lg('B')},"
                          f"\"<=\"&{chs}$C${k}),2),0)")
        for col, v in f.items():
            ws[f"{col}{r}"] = v
        arr(ws, f"AA{r}", (f"=IF(O{r}=0,\"\",MIN(IF(({lg('C')}=B{r})*({lg('F')}=\"OK\"),{lg('B')},99999)))"))


def build_calc_classes(ws):
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    labels = ["Valid", "Classroom", "Teacher", "Grade", "Grade order", "Students", "Participants",
              "Participation %", "Total $", "Avg / student", "Class goal", "% of class goal", "Key $",
              "Key part", "Rank $", "Rank part", "Filtered key $", "Filtered key part", "Grade sort"]
    labels += [f"Metric ch{i + 1}" for i in range(C.N_CH)] + ["Selected challenge key"]
    header_row(ws, 1, labels)
    for r in range(2, CLS_LAST + 1):
        f = {
            "A": f"=IF(AND(TRIM('Classrooms'!A{r})<>\"\",'Classrooms'!B{r}&\"\"<>\"\"),1,0)",
            "B": f"=IF(A{r}=1,TRIM('Classrooms'!A{r})&\" – \"&TRIM('Classrooms'!B{r}&\"\"),\"\")",
            "C": f"=IF(A{r}=1,TRIM('Classrooms'!A{r}),\"\")",
            "D": f"=IF(A{r}=1,TRIM('Classrooms'!B{r}&\"\"),\"\")",
            "E": f"=IF(A{r}=1,IFERROR(MATCH(D{r},GradeList,0),99),99)",
            "F": f"=IF(A{r}=1,COUNTIFS({cs('G')},B{r},{cs('A')},1),0)",
            "G": f"=IF(A{r}=1,COUNTIFS({cs('G')},B{r},{cs('P')},1),0)",
            "H": f"=IF(F{r}>0,ROUND(G{r}/F{r},6),0)",
            "I": f"=IF(A{r}=1,ROUND(SUMIFS({cs('N')},{cs('G')},B{r}),2),0)",
            "J": f"=IF(F{r}>0,I{r}/F{r},0)",
            "K": f"=IF(A{r}=1,N('Classrooms'!D{r}),0)",
            "L": f"=IF(K{r}>0,I{r}/K{r},\"\")",
            "M": f"=IF(A{r}=1,I{r}+(100-ROW())/10^7,-1)",
            "N": f"=IF(A{r}=1,H{r}+I{r}/10^10+(100-ROW())/10^14,-1)",
            "O": f"=IF(A{r}=1,COUNTIFS($A$2:$A${CLS_LAST},1,$I$2:$I${CLS_LAST},\">\"&I{r})+1,\"\")",
            "P": f"=IF(A{r}=1,COUNTIFS($A$2:$A${CLS_LAST},1,$H$2:$H${CLS_LAST},\">\"&H{r})+1,\"\")",
            "Q": f"=IF(AND(A{r}=1,OR(LB_Grade&\"\"=\"All grades\",D{r}=LB_Grade&\"\")),M{r},-1)",
            "R": f"=IF(AND(A{r}=1,OR(LB_Grade&\"\"=\"All grades\",D{r}=LB_Grade&\"\")),N{r},-1)",
            "S": f"=IF(A{r}=1,E{r}*100+ROW(),\"\")",
        }
        for i in range(C.N_CH):
            k = i + 2
            t = f"'Calc_Challenges'!$F${k}"
            mn = f"'Calc_Challenges'!$G${k}"
            f[CMET[i]] = (
                f"=IF(A{r}<>1,-1,IF({t}=\"CLASS_TOTAL_PER_GRADE\",ROUND(SUMIFS({cs(WIN[i])},{cs('G')},B{r}),2),"
                f"IF({t}=\"CLASS_PARTICIPATION_MIN\",IF(F{r}>0,ROUND(COUNTIFS({cs('G')},B{r},{cs('A')},1,"
                f"{cs(WIN[i])},\">=\"&MAX({mn},0.01))/F{r},6),0),"
                f"IF({t}=\"CLASS_PARTICIPATION_CUM\",IF(F{r}>0,ROUND(COUNTIFS({cs('G')},B{r},{cs('A')},1,"
                f"{cs(ASOF[i])},\">=\"&MAX({mn},PartMin,0.01))/F{r},6),0),-1))))")
        sel = f"INDEX(${CMET[0]}{r}:${CMET[-1]}{r},1,CH_Sel)"
        f["AI"] = f"=IF(A{r}<>1,-1,IF({sel}<0,-1,{sel}+(100-ROW())/10^10))"
        for col, v in f.items():
            ws[f"{col}{r}"] = v


def build_calc_grades(ws):
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    labels = ["Grade", "Order", "Classes", "Students", "Participants", "Participation %", "This year",
              "Avg / student", "Key", "Rank", "Last year", "Change vs last year", "Grade label"]
    labels += [f"Window $ ch{i + 1}" for i in range(C.N_CH)] + ["Selected challenge key"]
    header_row(ws, 1, labels)
    for i, g in enumerate(C.GRADES):
        r = 2 + i
        ws[f"A{r}"] = g
        ws[f"A{r}"].number_format = "@"
        f = {
            "B": i + 1,
            "C": f"=COUNTIFS({ccl('D')},A{r},{ccl('A')},1)",
            "D": f"=COUNTIFS({cs('E')},A{r},{cs('A')},1)",
            "E": f"=COUNTIFS({cs('E')},A{r},{cs('P')},1)",
            "F": f"=IF(D{r}>0,ROUND(E{r}/D{r},6),0)",
            "G": f"=ROUND(SUMIFS({cs('N')},{cs('E')},A{r},{cs('A')},1),2)",
            "H": f"=IF(D{r}>0,G{r}/D{r},0)",
            "I": f"=G{r}+(100-ROW())/10^7",
            "J": f"=COUNTIF($G$2:$G${G_LAST},\">\"&G{r})+1",
            "K": (f"=IFERROR(N(INDEX('History'!$F$4:$P$13,MATCH(CompYear,'History'!$A$4:$A$13,0),B{r})),0)"),
            "L": f"=G{r}-K{r}",
            "M": f"=IF(LEN(A{r})>1,A{r},\"Gr \"&A{r})",
        }
        for j in range(C.N_CH):
            t = f"'Calc_Challenges'!$F${j + 2}"
            f[GMET[j]] = f"=IF({t}=\"GRADE_TOTAL\",ROUND(SUMIFS({cs(WIN[j])},{cs('E')},A{r}),2),-1)"
        sel = f"INDEX(${GMET[0]}{r}:${GMET[-1]}{r},1,CH_Sel)"
        f["AC"] = f"=IF({sel}<0,-1,{sel}+(100-ROW())/10^7)"
        for col, v in f.items():
            ws[f"{col}{r}"] = v


def build_calc_families(ws):
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    header_row(ws, 1, ["#", "Family", "Students", "Participants", "Total $", "Key", "Rank", "Display"])
    for r in range(2, FAM_LAST + 1):
        ws[f"A{r}"] = r - 1
        ws[f"B{r}"] = f"=IFERROR(INDEX({cs('H')},MATCH(A{r},{cs('AF')},0)),\"\")"
        ws[f"C{r}"] = f"=IF(B{r}=\"\",0,COUNTIFS({cs('H')},B{r},{cs('A')},1))"
        ws[f"D{r}"] = f"=IF(B{r}=\"\",0,COUNTIFS({cs('H')},B{r},{cs('P')},1))"
        ws[f"E{r}"] = f"=IF(B{r}=\"\",0,ROUND(SUMIFS({cs('N')},{cs('H')},B{r}),2))"
        ws[f"F{r}"] = f"=IF(B{r}=\"\",-1,E{r}+(1000-ROW())/10^7)"
        ws[f"G{r}"] = f"=IF(B{r}=\"\",\"\",COUNTIF($E$2:$E${FAM_LAST},\">\"&E{r})+1)"
        ws[f"H{r}"] = (f"=IF(B{r}=\"\",\"\",IF(Privacy=\"Initials only\",LEFT(B{r},1)&\". Family\","
                       f"\"The \"&B{r}&\" Family\"))")


def build_calc_donors(ws):
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    header_row(ws, 1, ["#", "Donor", "Key", "Gifts", "Total $", "Sort key", "Students supported (rows)"])
    for r in range(2, DON_LAST + 1):
        ws[f"A{r}"] = r - 1
        ws[f"C{r}"] = f"=IFERROR(INDEX({lg('G')},MATCH(A{r},{lg('I')},0)),\"\")"
        ws[f"B{r}"] = f"=IF(C{r}=\"\",\"\",INDEX({lg('D')},MATCH(A{r},{lg('I')},0)))"
        ws[f"D{r}"] = f"=IF(C{r}=\"\",0,COUNTIFS({lg('G')},C{r}))"
        ws[f"E{r}"] = f"=IF(C{r}=\"\",0,ROUND(SUMIFS({lg('E')},{lg('G')},C{r}),2))"
        ws[f"F{r}"] = f"=IF(C{r}=\"\",-1,E{r}+(10000-ROW())/10^8)"


def build_calc_daily(ws):
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    header_row(ws, 1, ["Day #", "Date", "Weekday", "Online $", "Offline $", "Total $", "Gifts", "This year",
                       "Participants to date", "Last year", "This year", "Goal", "Label", "",
                       "Item", "Value"])
    for d in range(1, C.N_DAY + 1):
        r = d + 1
        ws[f"A{r}"] = d
        ws[f"B{r}"] = f"=IF(StartDate+A{r}-1<=EndDate,StartDate+A{r}-1,\"\")"
        ws[f"C{r}"] = f"=IF(B{r}=\"\",\"\",TEXT(B{r},\"ddd\"))"
        ws[f"D{r}"] = f"=IF(B{r}=\"\",\"\",ROUND(SUMIFS({lg('E')},{lg('B')},B{r},{lg('A')},\"Online\"),2))"
        ws[f"E{r}"] = f"=IF(B{r}=\"\",\"\",ROUND(SUMIFS({lg('E')},{lg('B')},B{r},{lg('A')},\"<>Online\"),2))"
        ws[f"F{r}"] = f"=IF(B{r}=\"\",\"\",D{r}+E{r})"
        ws[f"G{r}"] = f"=IF(B{r}=\"\",\"\",COUNTIFS({lg('B')},B{r},{lg('F')},\"OK\"))"
        ws[f"H{r}"] = f"=IF(B{r}=\"\",\"\",ROUND($P$2+SUM($F$2:F{r}),2))"
        ws[f"I{r}"] = f"=IF(B{r}=\"\",\"\",COUNTIFS({cs('AA')},\"<=\"&B{r},{cs('P')},1))"
        ly = "INDEX('History'!$B$17:$F$47,A{r},MATCH(CompYear,'History'!$B$16:$F$16,0))".format(r=r)
        ws[f"J{r}"] = f"=IF(B{r}=\"\",NA(),IFERROR(IF({ly}=\"\",NA(),{ly}),NA()))"
        ws[f"K{r}"] = f"=IF(B{r}=\"\",NA(),IF(B{r}>AsOfDate,NA(),H{r}))"
        ws[f"L{r}"] = f"=IF(B{r}=\"\",NA(),Goal)"
        ws[f"M{r}"] = f"=IF(B{r}=\"\",\"\",TEXT(B{r},\"m/d\"))"
    ws["O2"], ws["P2"] = "Before campaign start", (f"=ROUND(SUMIFS({lg('E')},{lg('B')},\"<\"&StartDate),2)")
    ws["O3"], ws["P3"] = "After campaign end", (f"=ROUND(SUMIFS({lg('E')},{lg('B')},\">\"&EndDate),2)")
    ws["O4"], ws["P4"] = "Campaign days", "=COUNT($B$2:$B$32)"


def build_calc_challenges(ws):
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    header_row(ws, 1, ["#", "Start", "End", "Name", "Reward", "Type", "Min $", "Target", "Needs window",
                       "Needs as-of", "Status", "Class max", "Class ties", "Class #1", "Class #2", "Grade max",
                       "Grade ties", "Grade #1", "Grade #2", "School value", "Hit", "Drawing eligible",
                       "Grade winners", "Class winners", "Headline", "Picker label", "Day", "Description"])
    for i in range(C.N_CH):
        r = i + 2
        cm = ccl(CMET[i])
        gm = cg(GMET[i])
        gba = rng("Calc_Challenges", BA[i], GB_ROW0, GB_LAST)
        cba = rng("Calc_Challenges", BA[i], CB_ROW0, CB_LAST)
        f = {
            "A": i + 1,
            "B": f"=IF(ISNUMBER('Challenges'!A{r}),INT('Challenges'!A{r}),\"\")",
            "C": f"=IF(B{r}=\"\",\"\",IF(ISNUMBER('Challenges'!B{r}),INT('Challenges'!B{r}),B{r}))",
            "D": f"=IF(B{r}=\"\",\"\",'Challenges'!D{r}&\"\")",
            "E": f"=IF(B{r}=\"\",\"\",'Challenges'!F{r}&\"\")",
            "F": f"=IF(B{r}=\"\",\"NONE\",IF(TRIM('Challenges'!G{r})=\"\",\"NONE\",UPPER(TRIM('Challenges'!G{r}))))",
            "G": f"=N('Challenges'!H{r})",
            "H": f"=N('Challenges'!I{r})",
            "I": (f"=IF(OR(F{r}=\"CLASS_TOTAL_PER_GRADE\",F{r}=\"CLASS_PARTICIPATION_MIN\","
                  f"F{r}=\"TOP_STUDENT_PER_CLASS\",F{r}=\"GRADE_TOTAL\",F{r}=\"DRAWING\"),1,0)"),
            "J": (f"=IF(OR(F{r}=\"CLASS_PARTICIPATION_CUM\",F{r}=\"SCHOOL_PARTICIPATION_TARGET\","
                  f"F{r}=\"SCHOOL_TOTAL_TARGET\"),1,0)"),
            "K": f"=IF(B{r}=\"\",\"\",IF(AsOfDate<B{r},\"Upcoming\",IF(AsOfDate<=C{r},\"LIVE\",\"Final\")))",
            "L": f"=IF(OR(F{r}=\"CLASS_PARTICIPATION_MIN\",F{r}=\"CLASS_PARTICIPATION_CUM\"),MAX({cm}),\"\")",
            "M": f"=IF(L{r}=\"\",\"\",COUNTIF({cm},L{r}))",
            "N": f"=IF(L{r}=\"\",\"\",INDEX({ccl('B')},MATCH(L{r},{cm},0)))",
            "P": f"=IF(F{r}=\"GRADE_TOTAL\",MAX({gm}),\"\")",
            "Q": f"=IF(P{r}=\"\",\"\",COUNTIF({gm},P{r}))",
            "R": f"=IF(P{r}=\"\",\"\",INDEX({cg('A')},MATCH(P{r},{gm},0)))",
            "T": (f"=IF(F{r}=\"SCHOOL_PARTICIPATION_TARGET\",IF(Students=0,0,ROUND(COUNTIFS({cs('A')},1,"
                  f"{cs(ASOF[i])},\">=\"&MAX(G{r},PartMin,0.01))/Students,6)),IF(F{r}=\"SCHOOL_TOTAL_TARGET\","
                  f"ROUND(SUM({cs(ASOF[i])}),2),\"\"))"),
            "U": f"=IF(T{r}=\"\",\"\",IF(T{r}>=H{r},\"HIT\",IF(K{r}=\"Final\",\"MISSED\",\"NOT YET\")))",
            "V": f"=IF(F{r}=\"DRAWING\",COUNTIFS({cs('A')},1,{cs(WIN[i])},\">0\"),\"\")",
            "W": f"=IF(F{r}=\"CLASS_TOTAL_PER_GRADE\",COUNTIF({gba},\">0\"),\"\")",
            "X": f"=IF(F{r}=\"TOP_STUDENT_PER_CLASS\",COUNTIF({cba},\">0\"),\"\")",
            "Y": (f"=IF(B{r}=\"\",\"\",IF(F{r}=\"CLASS_TOTAL_PER_GRADE\",IF(W{r}=0,\"No gifts yet\","
                  f"W{r}&\" grade winners (see Challenge Detail)\"),"
                  f"IF(F{r}=\"TOP_STUDENT_PER_CLASS\",IF(X{r}=0,\"No gifts yet\","
                  f"X{r}&\" class winners (see Challenge Detail)\"),"
                  f"IF(OR(F{r}=\"CLASS_PARTICIPATION_MIN\",F{r}=\"CLASS_PARTICIPATION_CUM\"),IF(L{r}<=0,"
                  f"\"No gifts yet\",IF(M{r}>1,\"TIE: \"&N{r}&\" / \"&O{r}&IF(M{r}>2,\" +\"&(M{r}-2)&\" more\",\"\"),"
                  f"N{r})&\": \"&TEXT(L{r},\"0%\")),"
                  f"IF(F{r}=\"GRADE_TOTAL\",IF(P{r}<=0,\"No gifts yet\",IF(Q{r}>1,\"TIE: Grades \"&R{r}&\" / \"&S{r},"
                  f"\"Grade \"&R{r})&\": \"&TEXT(P{r},\"$#,##0\")),"
                  f"IF(F{r}=\"SCHOOL_PARTICIPATION_TARGET\",U{r}&\": \"&TEXT(T{r},\"0%\")&\" of \"&TEXT(H{r},\"0%\")"
                  f"&\" goal\",IF(F{r}=\"SCHOOL_TOTAL_TARGET\",U{r}&\": \"&TEXT(T{r},\"$#,##0\")&\" of \"&"
                  f"TEXT(H{r},\"$#,##0\"),IF(F{r}=\"DRAWING\",V{r}&\" eligible; draw \"&H{r}&\" winners in "
                  f"Challenge Detail\",AB{r}))))))))"),
            "Z": f"=IF(B{r}=\"\",\"\",A{r}&\". \"&D{r}&\" (\"&TEXT(B{r},\"m/d\")&\")\")",
            "AA": f"=IF(B{r}=\"\",\"\",'Challenges'!C{r}&\"\")",
            "AB": f"=IF(B{r}=\"\",\"\",'Challenges'!E{r}&\"\")",
        }
        for col, v in f.items():
            ws[f"{col}{r}"] = v
        arr(ws, f"O{r}", (f"=IF(L{r}=\"\",\"\",IF(M{r}<2,\"\",INDEX({ccl('B')},SMALL(IF({cm}=L{r},"
                          f"ROW({cm})-1,9999),2))))"))
        arr(ws, f"S{r}", (f"=IF(P{r}=\"\",\"\",IF(Q{r}<2,\"\",INDEX({cg('A')},SMALL(IF({gm}=P{r},"
                          f"ROW({gm})-1,99),2))))"))

    # Per-grade winners block (CLASS_TOTAL_PER_GRADE)
    put(ws, f"A{GB_ROW0 - 1}", "Per-grade class winners", font=S.font(11, True))
    for gi, g in enumerate(C.GRADES):
        r = GB_ROW0 + gi
        ws[f"A{r}"] = g
        ws[f"A{r}"].number_format = "@"
        for i in range(C.N_CH):
            k = i + 2
            cm = ccl(CMET[i])
            a = f"{BA[i]}{r}"
            arr(ws, a, (f"=IF($F${k}<>\"CLASS_TOTAL_PER_GRADE\",\"\",MAX(IF({ccl('D')}=$A{r},{cm},-1)))"))
            cnt = f"COUNTIFS({ccl('D')},$A{r},{cm},{a})"
            pick = f"INDEX({ccl('B')},SMALL(IF(({ccl('D')}=$A{r})*({cm}={a}),ROW({cm})-1,9999),{{n}}))"
            arr(ws, f"{BW[i]}{r}", (f"=IF({a}=\"\",\"\",IF({a}<=0,\"—\",IF({cnt}>1,\"TIE: \",\"\")&"
                                    f"{pick.format(n=1)}&IF({cnt}>1,\" / \"&{pick.format(n=2)},\"\")))"))
    # Per-class top student block (TOP_STUDENT_PER_CLASS)
    put(ws, f"A{CB_ROW0 - 1}", "Per-class top students", font=S.font(11, True))
    for ci in range(C.N_CLS):
        r = CB_ROW0 + ci
        ws[f"A{r}"] = f"='Calc_Classes'!B{ci + 2}"
        for i in range(C.N_CH):
            k = i + 2
            w = cs(WIN[i])
            a = f"{BA[i]}{r}"
            arr(ws, a, (f"=IF(OR($F${k}<>\"TOP_STUDENT_PER_CLASS\",$A{r}=\"\"),\"\","
                        f"MAX(IF({cs('G')}=$A{r},{w},-1)))"))
            cnt = f"COUNTIFS({cs('G')},$A{r},{w},{a})"
            pick = f"INDEX({cs('AG')},SMALL(IF(({cs('G')}=$A{r})*({w}={a}),ROW({w})-1,99999),{{n}}))"
            arr(ws, f"{BW[i]}{r}", (f"=IF({a}=\"\",\"\",IF({a}<=0,\"—\",IF({cnt}>1,\"TIE: \",\"\")&"
                                    f"{pick.format(n=1)}&IF({cnt}>1,\" / \"&{pick.format(n=2)},\"\")))"))


def build_calc_kpi(wb, ws):
    ws.sheet_properties.tabColor = S.TAB_ENGINE
    widths(ws, {"A": 28, "B": 60})
    items = [
        ("AsOfDate", "As-of date", "=IF(ISNUMBER('Settings'!B15),INT('Settings'!B15),TODAY())"),
        ("Raised", "Raised", f"=ROUND(SUM({cs('N')}),2)"),
        ("RaisedOnline", "Online", f"=ROUND(SUM({cs('L')}),2)"),
        ("RaisedOffline", "Cash/check", f"=ROUND(SUM({cs('M')}),2)"),
        ("Remaining", "Remaining", "=MAX(0,Goal-Raised)"),
        ("PctGoal", "% of goal", "=IF(Goal>0,Raised/Goal,0)"),
        ("Students", "Students", f"=COUNTIF({cs('A')},1)"),
        ("Participants", "Participants", f"=SUM({cs('P')})"),
        ("PartPct", "Participation %", "=IF(Students>0,Participants/Students,0)"),
        ("Gifts", "Gifts", f"=SUM({cs('O')})"),
        ("DaysLeft", "Days left", "=MAX(0,EndDate-AsOfDate)"),
        ("RaisedToday", "Raised on as-of date", f"=ROUND(SUMIFS({lg('E')},{lg('B')},AsOfDate),2)"),
        ("AvgStudent", "Avg per student", "=IF(Students>0,Raised/Students,0)"),
        ("Families", "Families", f"=MAX({cs('AF')})"),
        ("Donors", "Unique donors", f"=MAX({lg('I')})"),
        ("ThermoCells", "Thermometer cells lit (of 40)", "=MIN(40,MAX(0,ROUND(PctGoal*40,0)))"),
        ("StretchCell", "Stretch marker cell", "=MIN(40,MAX(1,ROUND(Stretch/Goal*40,0)))"),
        ("TodayIdx", "Challenge index today", None),
        ("NextIdx", "Next challenge index", None),
        ("ShowIdx", "Challenge shown on TV", "=IF(TodayIdx<999,TodayIdx,IF(NextIdx<999,NextIdx,0))"),
        ("ShowLabel", "TV challenge label",
         "=IF(TodayIdx<999,\"TODAY'S CHALLENGE\",IF(NextIdx<999,\"NEXT CHALLENGE: \"&TEXT(INDEX("
         + ch('B') + ",NextIdx),\"ddd m/d\"),\"CHALLENGES COMPLETE\"))"),
        ("ShowName", "TV challenge name", f"=IF(ShowIdx=0,\"Thank you, Saints!\",INDEX({ch('D')},ShowIdx))"),
        ("ShowReward", "TV challenge reward", f"=IF(ShowIdx=0,\"\",INDEX({ch('E')},ShowIdx))"),
        ("ShowHeadline", "TV challenge standing",
         f"=IF(ShowIdx=0,\"\",IF(TodayIdx<999,INDEX({ch('Y')},ShowIdx),INDEX({ch('AB')},ShowIdx)))"),
        ("LYTotal", "Last year total",
         "=IFERROR(N(INDEX('History'!$C$4:$C$13,MATCH(CompYear,'History'!$A$4:$A$13,0))),0)"),
        ("LYSameDay", "Last year same campaign day",
         "=IFERROR(N(INDEX('History'!$B$17:$F$47,MAX(1,MIN(31,AsOfDate-StartDate+1)),"
         "MATCH(CompYear,'History'!$B$16:$F$16,0))),0)"),
        ("Banner", "TV banner",
         "=IF(Raised>=Goal,\"GOAL REACHED! THANK YOU, SAINTS!\",IF(Raised>=Stretch,\"WE HIT \"&TEXT(Stretch,"
         "\"$#,##0\")&\": \"&StretchLabel,TEXT(Stretch-Raised,\"$#,##0\")&\" to go until \"&StretchLabel))"),
        ("CH_Sel", "Challenge Detail selection",
         f"=IFERROR(MATCH('Challenge Winners'!$C$25,{ch('Z')},0),1)"),
        ("CH_SelType", "Selected challenge type", f"=INDEX({ch('F')},CH_Sel)"),
        ("PR_SelLevel", "Prize Report selection",
         f"=IFERROR(MATCH('Prize Report'!$C$4,'Prizes'!$A$2:$A${PR_LAST},0),0)"),
        ("PrizeCount", "Prize tiers", f"=COUNTIFS('Prizes'!$B$2:$B${PR_LAST},\">0\")"),
        ("ThermoMax", "Thermometer label", "=TEXT(PctGoal,\"0%\")"),
        ("IssueCount", "Data Check issues", "='Data Check'!$C$4"),
    ]
    header_row(ws, 1, ["Name", "Value"])
    for i, (nm, label, formula) in enumerate(items):
        r = 2 + i
        ws[f"A{r}"] = f"{nm}  ({label})"
        if formula:
            ws[f"B{r}"] = formula
        name(wb, nm, f"'Calc_KPI'!$B${r}")
        if nm == "TodayIdx":
            arr(ws, f"B{r}", (f"=MIN(IF(ISNUMBER({ch('B')}),IF(({ch('B')}<=AsOfDate)*({ch('C')}>=AsOfDate),"
                              f"{ch('A')},999),999))"))
        if nm == "NextIdx":
            arr(ws, f"B{r}", f"=MIN(IF(ISNUMBER({ch('B')}),IF({ch('B')}>AsOfDate,{ch('A')},999),999))")


# ----------------------------------------------------------------------------
# OUTPUT SHEETS
# ----------------------------------------------------------------------------
def build_tv(ws, days):
    ws.sheet_properties.tabColor = S.TAB_OUTPUT
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False
    ws.sheet_view.zoomScale = 80
    widths(ws, {"A": 3, "B": 9, "C": 8, "D": 26, "E": 2})
    for col in "FGHIJKLM":
        ws.column_dimensions[col].width = 12
    widths(ws, {"N": 2, "O": 5, "P": 22, "Q": 22, "R": 12, "S": 9, "T": 2})
    for r in range(1, 60):
        for col in [CL(i) for i in range(1, 21)]:
            ws[f"{col}{r}"].fill = S.fill(S.CREAM)

    ws.merge_cells("B2:S2")
    put(ws, "B2", "=UPPER(CampaignName)&\"  \"&CampaignYear", font=S.display(40, S.RED), align=S.CENTER)
    ws.row_dimensions[2].height = 56
    ws.merge_cells("B3:S3")
    put(ws, "B3", "=SchoolName&\"  •  \"&Tagline", font=S.display(16, S.NAVY), align=S.CENTER)
    ws.row_dimensions[3].height = 26
    ws.merge_cells("B4:S4")
    put(ws, "B4", "=Banner", font=S.display(14, S.WHITE), fill=S.GREEN, align=S.CENTER)
    ws.row_dimensions[4].height = 26

    # Thermometer: 40 cells (rows 6-45) + bulb (rows 46-48). Col A holds the cell index (hidden by color).
    put(ws, "C5", "GOAL", font=S.display(10, S.NAVY), align=S.CENTER)
    for i in range(40):
        r = 6 + i
        idx = 40 - i
        ws.row_dimensions[r].height = 14
        put(ws, f"A{r}", idx, font=S.font(6, color=S.CREAM))
        put(ws, f"B{r}", f"=IF(MOD(A{r},4)=0,TEXT(Goal*A{r}/40,\"$#,##0\"),\"\")",
            font=S.font(9, True, S.NAVY), align=S.RIGHT)
        c = put(ws, f"C{r}", fill="E8E8E8")
        c.border = S.Border(left=S.thick_navy, right=S.thick_navy, top=S.thick_navy if idx == 40 else None)
        put(ws, f"D{r}", (f"=IF(A{r}=40,\"◄ GOAL \"&TEXT(Goal,\"$#,##0\"),IF(A{r}=StretchCell,\"◄ \"&"
                          f"TEXT(Stretch,\"$#,##0\")&\" SLIME!\",\"\"))"), font=S.font(10, True, S.RED))
    ws.conditional_formatting.add("C6:C45", FormulaRule(formula=["$A6<=ThermoCells"], fill=S.fill(S.RED)))
    ws.merge_cells("B46:D48")
    put(ws, "B46", "=TEXT(PctGoal,\"0%\")&\" of goal\"", font=S.display(20, S.WHITE), fill=S.RED, align=S.CENTER)
    for r in (46, 47, 48):
        ws.row_dimensions[r].height = 18

    # KPI tiles
    def tile(cells_label, cells_value, label, formula, fmt, size, fill_color):
        ws.merge_cells(cells_label)
        put(ws, cells_label.split(":")[0], label, font=S.font(12, True, S.WHITE), fill=fill_color, align=S.CENTER)
        ws.merge_cells(cells_value)
        put(ws, cells_value.split(":")[0], formula, font=S.display(size, fill_color), fill=S.WHITE,
            fmt=fmt, align=S.CENTER)

    tile("F6:I7", "F8:I13", "RAISED SO FAR", "=Raised", S.FMT_USD, 40, S.RED)
    tile("J6:M7", "J8:M13", "% OF GOAL", "=PctGoal", S.FMT_PCT, 40, S.NAVY)
    ws.merge_cells("F14:M15")
    put(ws, "F14", "=\"Goal \"&TEXT(Goal,\"$#,##0\")&\"   •   \"&TEXT(Remaining,\"$#,##0\")&\" to go   •   \"&"
        "TEXT(Gifts,\"#,##0\")&\" gifts\"", font=S.font(13, True, S.NAVY), align=S.CENTER)
    tile("F17:G18", "F19:G22", "PARTICIPATION", "=PartPct", S.FMT_PCT, 22, S.GREEN)
    tile("H17:I18", "H19:I22", "RAISED TODAY", "=RaisedToday", S.FMT_USD, 20, S.RED)
    tile("J17:K18", "J19:K22", "DAYS LEFT", "=DaysLeft", "0", 22, S.NAVY)
    tile("L17:M18", "L19:M22", "LAST YEAR, SAME DAY", "=LYSameDay", S.FMT_USD, 18, "6A6A6A")

    ws.merge_cells("F24:M25")
    put(ws, "F24", "=ShowLabel", font=S.display(14, S.WHITE), fill=S.NAVY, align=S.CENTER)
    ws.merge_cells("F26:M28")
    put(ws, "F26", "=ShowName", font=S.display(18, S.RED), fill=S.SAND, align=S.CENTER)
    ws.merge_cells("F29:M30")
    put(ws, "F29", "=ShowReward", font=S.font(13, True, S.NAVY), fill=S.SAND, align=S.CENTER)
    ws.merge_cells("F31:M32")
    put(ws, "F31", "=ShowHeadline", font=S.font(13, True, S.GREEN), fill=S.SAND, align=S.CENTER)

    # Leaderboards (privacy-safe display names)
    ws.merge_cells("O6:S6")
    put(ws, "O6", "TOP 10 CLASSROOMS", font=S.display(14, S.WHITE), fill=S.RED, align=S.CENTER)
    ws.row_dimensions[6].height = 22
    for k in range(1, 11):
        r = 6 + k
        v = f"LARGE({ccl('M')},{k})"
        hit = f"MATCH({v},{ccl('M')},0)"
        put(ws, f"O{r}", k, font=S.display(11, S.NAVY), align=S.CENTER)
        ws.merge_cells(f"P{r}:Q{r}")
        put(ws, f"P{r}", f"=IF({v}<0.005,\"\",INDEX({ccl('B')},{hit}))", font=S.font(12, True))
        put(ws, f"R{r}", f"=IF({v}<0.005,\"\",INDEX({ccl('I')},{hit}))", font=S.font(12, True), fmt=S.FMT_USD)
        put(ws, f"S{r}", f"=IF({v}<0.005,\"\",INDEX({ccl('H')},{hit}))", font=S.font(11, color=S.GREEN),
            fmt=S.FMT_PCT)
    ws.merge_cells("O18:S18")
    put(ws, "O18", "TOP 10 FUNDRAISERS", font=S.display(14, S.WHITE), fill=S.NAVY, align=S.CENTER)
    for k in range(1, 11):
        r = 18 + k
        v = f"LARGE({cs('Q')},{k})"
        hit = f"MATCH({v},{cs('Q')},0)"
        put(ws, f"O{r}", k, font=S.display(11, S.NAVY), align=S.CENTER)
        put(ws, f"P{r}", f"=IF({v}<0.005,\"\",INDEX({cs('S')},{hit}))", font=S.font(12, True))
        put(ws, f"Q{r}", f"=IF({v}<0.005,\"\",INDEX({cs('G')},{hit}))", font=S.font(10, color="555555"))
        ws.merge_cells(f"R{r}:S{r}")
        put(ws, f"R{r}", f"=IF({v}<0.005,\"\",INDEX({cs('N')},{hit}))", font=S.font(12, True), fmt=S.FMT_USD)

    # Charts
    gws = ws.parent["Calc_Grades"]
    bar = BarChart()
    bar.type, bar.style = "col", 10
    bar.title = "Raised by Grade (vs last year)"
    bar.y_axis.numFmt = '"$"#,##0'
    bar.y_axis.majorGridlines = None
    bar.add_data(Reference(gws, min_col=7, min_row=1, max_row=G_LAST), titles_from_data=True)
    bar.add_data(Reference(gws, min_col=11, min_row=1, max_row=G_LAST), titles_from_data=True)
    bar.set_categories(Reference(gws, min_col=13, min_row=2, max_row=G_LAST))
    bar.series[0].graphicalProperties.solidFill = S.RED
    bar.series[1].graphicalProperties.solidFill = "BFBFBF"
    bar.gapWidth = 40
    bar.legend.position = "b"
    bar.width, bar.height = 19.5, 8.2
    bar.visible_cells_only = False
    ws.add_chart(bar, "F34")

    dws = ws.parent["Calc_Daily"]
    line = LineChart()
    line.title = "Cumulative $ by Day"
    line.y_axis.numFmt = '"$"#,##0'
    line.y_axis.majorGridlines = None
    for col, color, dash in [(11, S.RED, None), (10, "9E9E9E", "dash"), (12, S.GREEN, "sysDot")]:
        line.add_data(Reference(dws, min_col=col, min_row=1, max_row=1 + days), titles_from_data=True)
        s = line.series[-1]
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 38000 if col == 11 else 20000
        s.smooth = False
        if dash:
            s.graphicalProperties.line.dashStyle = dash
    line.set_categories(Reference(dws, min_col=13, min_row=2, max_row=1 + days))
    line.legend.position = "b"
    line.width, line.height = 14.6, 9.6
    line.visible_cells_only = False
    ws.add_chart(line, "O30")

    ws.merge_cells("B50:S50")
    put(ws, "B50", "=\"Updated through \"&TEXT(AsOfDate,\"dddd, mmmm d\")&\"   •   Every gift counts. GO SAINTS!\"",
        font=S.font(11, italic=True, color="555555"), align=S.CENTER)
    ws.print_area = "A1:T50"
    print_setup(ws, landscape=True)
    ws.page_setup.fitToHeight = 1
    protect(ws)


def build_leaderboards(ws):
    ws.sheet_properties.tabColor = S.TAB_OUTPUT
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:Q1", "LEADERBOARDS", 22, fill=S.RED)
    put(ws, "A3", "Filter by grade:", font=S.font(12, True))
    put(ws, "C3", "All grades", fill=S.INPUT, font=S.font(12, True), border=S.BOX)
    add_list_dv(ws, "C3", '"All grades,' + ",".join(C.GRADES) + '"')
    unlock(ws, "C3")
    put(ws, "E3", "Names follow the privacy setting on the Settings tab.", font=S.font(9, italic=True))
    widths(ws, {"A": 5, "B": 22, "C": 24, "D": 12, "E": 2, "F": 5, "G": 28, "H": 12, "I": 10, "J": 2,
                "K": 5, "L": 28, "M": 10, "N": 12, "O": 2, "P": 5, "Q": 26, "R": 12})

    def table(col0, row0, title, cols, n, key_rng, fields, color):
        c0 = col0
        last = CL(ws[f"{col0}1"].column + len(cols) - 1)
        ws.merge_cells(f"{c0}{row0}:{last}{row0}")
        put(ws, f"{c0}{row0}", title, font=S.display(12, S.WHITE), fill=color, align=S.CENTER)
        header_row(ws, row0 + 1, cols, ws[f"{col0}1"].column, fill=S.NAVY, height=20)
        for k in range(1, n + 1):
            r = row0 + 1 + k
            v = f"LARGE({key_rng},{k})"
            hit = f"MATCH({v},{key_rng},0)"
            put(ws, f"{c0}{r}", k, align=S.CENTER, font=S.font(11, True, S.NAVY), border=S.BOX)
            for j, (src, fmt, thr) in enumerate(fields):
                col = CL(ws[f"{col0}1"].column + 1 + j)
                put(ws, f"{col}{r}", f"=IF({v}<{thr},\"\",INDEX({src},{hit}))", fmt=fmt, border=S.BOX)
            if k % 2 == 0:
                style_range(ws, f"{c0}{r}:{last}{r}", fill="F7F3EA")

    table("A", 5, "TOP STUDENTS", ["#", "Student", "Classroom", "Raised"], 10, cs("AB"),
          [(cs("S"), None, 0.005), (cs("G"), None, 0.005), (cs("N"), S.FMT_USD, 0.005)], S.RED)
    table("F", 5, "TOP CLASSROOMS ($)", ["#", "Classroom", "Raised", "Particip."], 10, ccl("Q"),
          [(ccl("B"), None, 0.005), (ccl("I"), S.FMT_USD, 0.005), (ccl("H"), S.FMT_PCT, 0.005)], S.NAVY)
    table("K", 5, "TOP CLASSROOMS (PARTICIPATION)", ["#", "Classroom", "Particip.", "Raised"], 10, ccl("R"),
          [(ccl("B"), None, 1e-9), (ccl("H"), S.FMT_PCT, 1e-9), (ccl("I"), S.FMT_USD, 1e-9)], S.GREEN)
    table("P", 5, "TOP FAMILIES", ["#", "Family", "Raised"], 10, rng("Calc_Families", "F", 2, FAM_LAST),
          [(rng("Calc_Families", "H", 2, FAM_LAST), None, 0.005),
           (rng("Calc_Families", "E", 2, FAM_LAST), S.FMT_USD, 0.005)], S.GOLD)
    # Grades: all 11, ranked
    ws.merge_cells("A19:I19")
    put(ws, "A19", "GRADES (ALL, RANKED BY $)", font=S.display(12, S.WHITE), fill=S.GREEN, align=S.CENTER)
    header_row(ws, 20, ["#", "Grade", "Classes", "Students", "Participants", "Participation", "Raised",
                        "Avg / student", "vs last year"], 1, height=20)
    for k in range(1, len(C.GRADES) + 1):
        r = 20 + k
        v = f"LARGE({cg('I')},{k})"
        hit = f"MATCH({v},{cg('I')},0)"
        put(ws, f"A{r}", k, align=S.CENTER, border=S.BOX)
        for col, src, fmt in [("B", "A", None), ("C", "C", "0"), ("D", "D", "0"), ("E", "E", "0"),
                              ("F", "F", S.FMT_PCT), ("G", "G", S.FMT_USD), ("H", "H", S.FMT_USD),
                              ("I", "L", '+"$"#,##0;-"$"#,##0;"$"0')]:
            put(ws, f"{col}{r}", f"=INDEX({cg(src)},{hit})", fmt=fmt, border=S.BOX)
        ws[f"B{r}"].value = f"=\"Grade \"&INDEX({cg('A')},{hit})"
    print_setup(ws, landscape=True)
    protect(ws)


def build_challenge_winners(ws, default_pick):
    ws.sheet_properties.tabColor = S.TAB_OUTPUT
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:H1", "DAILY CHALLENGE WINNERS", 22, fill=S.RED)
    put(ws, "A2", "=\"Results update automatically. As of \"&TEXT(AsOfDate,\"ddd m/d/yyyy\")",
        font=S.font(10, italic=True))
    widths(ws, {"A": 4, "B": 14, "C": 34, "D": 46, "E": 22, "F": 10, "G": 52, "H": 4, "I": 28, "J": 22})
    header_row(ws, 4, ["#", "Dates", "Challenge", "Reward", "Day", "Status", "Result / standing"], 1)
    for i in range(C.N_CH):
        r = 5 + i
        k = i + 2
        chs = "'Calc_Challenges'!"
        put(ws, f"A{r}", f"=IF({chs}B{k}=\"\",\"\",{chs}A{k})", align=S.CENTER, border=S.BOX)
        put(ws, f"B{r}", (f"=IF({chs}B{k}=\"\",\"\",TEXT({chs}B{k},\"m/d\")&IF({chs}C{k}>{chs}B{k},"
                          f"\"–\"&TEXT({chs}C{k},\"m/d\"),\"\"))"), border=S.BOX, align=S.CENTER)
        put(ws, f"C{r}", f"={chs}D{k}", border=S.BOX, font=S.font(11, True))
        put(ws, f"D{r}", f"={chs}E{k}", border=S.BOX)
        put(ws, f"E{r}", f"={chs}AA{k}", border=S.BOX)
        put(ws, f"F{r}", f"={chs}K{k}", border=S.BOX, align=S.CENTER)
        put(ws, f"G{r}", f"={chs}Y{k}", border=S.BOX, font=S.font(11, True, S.NAVY))
    ws.conditional_formatting.add("F5:F19", FormulaRule(formula=['F5="LIVE"'], fill=S.fill(S.GOLD),
                                                        font=S.font(11, True)))
    ws.conditional_formatting.add("G5:G19", FormulaRule(formula=['LEFT(G5,3)="HIT"'], fill=S.fill("D4EDDA")))

    put(ws, "A22", "CHALLENGE DETAIL", font=S.display(16, S.RED))
    put(ws, "A23", "Pick a challenge to see every winner (per grade / per class / ranking / drawing).",
        font=S.font(10, italic=True))
    put(ws, "B25", "Challenge:", font=S.font(12, True), align=S.RIGHT)
    ws.merge_cells("C25:D25")
    put(ws, "C25", default_pick, fill=S.INPUT, font=S.font(12, True), border=S.BOX)
    add_list_dv(ws, "C25", f"={ch('Z')}")
    unlock(ws, "C25:D25")
    put(ws, "E25", "=CH_SelType", font=S.font(10, italic=True, color="555555"))
    header_row(ws, 27, ["#", "Group", "Winner", "Classroom / note", "Value"], 1, height=22)
    t = "CH_SelType"
    for j in range(1, 41):
        r = 27 + j
        part_v = f"LARGE({ccl('AI')},{j})"
        grade_v = f"LARGE({cg('AC')},{j})"
        draw_v = f"LARGE({cs('AH')},{j})"
        draw_n = "MIN(INDEX(" + ch('H') + ",CH_Sel),INDEX(" + ch('V') + ",CH_Sel))"
        winner = (
            f"=IFERROR(IF({t}=\"CLASS_TOTAL_PER_GRADE\",IF({j}<={len(C.GRADES)},INDEX('Calc_Challenges'!"
            f"$B${GB_ROW0}:$AE${GB_LAST},{j},2*CH_Sel-1),\"\"),"
            f"IF({t}=\"TOP_STUDENT_PER_CLASS\",IF(INDEX('Calc_Challenges'!$A${CB_ROW0}:$A${CB_LAST},{j})=\"\",\"\","
            f"INDEX('Calc_Challenges'!$B${CB_ROW0}:$AE${CB_LAST},{j},2*CH_Sel-1)),"
            f"IF(OR({t}=\"CLASS_PARTICIPATION_MIN\",{t}=\"CLASS_PARTICIPATION_CUM\"),IF({part_v}<0,\"\","
            f"INDEX({ccl('B')},MATCH({part_v},{ccl('AI')},0))),"
            f"IF({t}=\"GRADE_TOTAL\",IF({grade_v}<0,\"\",\"Grade \"&INDEX({cg('A')},MATCH({grade_v},{cg('AC')},0))),"
            f"IF({t}=\"DRAWING\",IF({j}>{draw_n},\"\",INDEX({cs('AG')},MATCH({draw_v},{cs('AH')},0))),"
            f"IF(AND({j}=1,OR({t}=\"SCHOOL_PARTICIPATION_TARGET\",{t}=\"SCHOOL_TOTAL_TARGET\")),"
            f"INDEX({ch('U')},CH_Sel),\"\")))))),\"\")")
        group = (f"=IF(C{r}=\"\",\"\",IF({t}=\"CLASS_TOTAL_PER_GRADE\",\"Grade \"&INDEX({cg('A')},{j}),"
                 f"IF({t}=\"TOP_STUDENT_PER_CLASS\",INDEX('Calc_Challenges'!$A${CB_ROW0}:$A${CB_LAST},{j}),"
                 f"IF({t}=\"DRAWING\",\"Wild Card winner\",IF(LEFT({t},6)=\"SCHOOL\",\"Result\",\"Rank\")))))")
        note = (f"=IF(C{r}=\"\",\"\",IF({t}=\"DRAWING\",INDEX({cs('G')},MATCH({draw_v},{cs('AH')},0)),"
                f"IF({t}=\"SCHOOL_PARTICIPATION_TARGET\",\"Target \"&TEXT(INDEX({ch('H')},CH_Sel),\"0%\"),"
                f"IF({t}=\"SCHOOL_TOTAL_TARGET\",\"Target \"&TEXT(INDEX({ch('H')},CH_Sel),\"$#,##0\"),\"\"))))")
        value = (
            f"=IFERROR(IF(C{r}=\"\",\"\",IF({t}=\"CLASS_TOTAL_PER_GRADE\",TEXT(INDEX('Calc_Challenges'!"
            f"$B${GB_ROW0}:$AE${GB_LAST},{j},2*CH_Sel),\"$#,##0.00\"),"
            f"IF({t}=\"TOP_STUDENT_PER_CLASS\",TEXT(INDEX('Calc_Challenges'!$B${CB_ROW0}:$AE${CB_LAST},{j},"
            f"2*CH_Sel),\"$#,##0.00\"),"
            f"IF(OR({t}=\"CLASS_PARTICIPATION_MIN\",{t}=\"CLASS_PARTICIPATION_CUM\"),TEXT({part_v},\"0.0%\"),"
            f"IF({t}=\"GRADE_TOTAL\",TEXT({grade_v},\"$#,##0.00\"),"
            f"IF({t}=\"SCHOOL_PARTICIPATION_TARGET\",TEXT(INDEX({ch('T')},CH_Sel),\"0.0%\"),"
            f"IF({t}=\"SCHOOL_TOTAL_TARGET\",TEXT(INDEX({ch('T')},CH_Sel),\"$#,##0\"),\"\""
            f"))))))),\"\")")
        put(ws, f"A{r}", f"=IF(C{r}=\"\",\"\",{j})", align=S.CENTER)
        put(ws, f"B{r}", group)
        put(ws, f"C{r}", winner, font=S.font(11, True))
        put(ws, f"D{r}", note)
        put(ws, f"E{r}", value, align=S.RIGHT)
    row_lines(ws, "A28:E67", "$C28")
    put(ws, "G27", "Wild Card drawing tip", font=S.font(11, True, S.RED))
    tips = ["The drawing re-shuffles every time the workbook recalculates.",
            "To lock winners: select C28:D37, Copy, then Paste Values into I28.",
            "Locked winners (paste here):"]
    for i, tip in enumerate(tips):
        put(ws, f"G{28 + i}", tip, font=S.font(10, italic=i < 2))
    for r in range(28, 68):
        put(ws, f"I{r}", fill=S.INPUT, border=S.BOX)
        put(ws, f"J{r}", fill=S.INPUT, border=S.BOX)
    unlock(ws, "I28:J67")
    print_setup(ws, landscape=True)
    ws.print_area = "A1:G67"
    protect(ws)


def build_prize_report(ws, default_prize):
    ws.sheet_properties.tabColor = S.TAB_OUTPUT
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:F1", "PRIZE REPORT", 22, fill=S.RED)
    put(ws, "A2", "=\"Students who earned each milestone prize, as of \"&TEXT(AsOfDate,\"m/d/yyyy\")",
        font=S.font(10, italic=True))
    widths(ws, {"A": 6, "B": 30, "C": 30, "D": 14, "E": 26, "F": 14, "H": 6})
    put(ws, "B4", "Show students who earned:", font=S.font(12, True), align=S.RIGHT)
    put(ws, "C4", default_prize, fill=S.INPUT, font=S.font(12, True), border=S.BOX)
    add_list_dv(ws, "C4", f"='Prizes'!$A$2:$A${PR_LAST}")
    unlock(ws, "C4")
    put(ws, "D4", f"=IF(PR_SelLevel=0,\"\",\"(raised \"&TEXT(INDEX('Prizes'!$B$2:$B${PR_LAST},PR_SelLevel),"
        f"\"$#,##0\")&\"+)\")", font=S.font(11, italic=True))
    header_row(ws, 6, ["Tier", "Prize", "Raise at least", "Students", "", ""], 1, height=20)
    for i in range(C.N_PR):
        r = 7 + i
        p = i + 2
        put(ws, f"A{r}", f"=IF('Prizes'!A{p}=\"\",\"\",{i + 1})", align=S.CENTER)
        put(ws, f"B{r}", f"=IF('Prizes'!A{p}=\"\",\"\",'Prizes'!A{p})")
        put(ws, f"C{r}", f"=IF('Prizes'!A{p}=\"\",\"\",'Prizes'!B{p})", fmt=S.FMT_USD)
        put(ws, f"D{r}", f"=IF('Prizes'!A{p}=\"\",\"\",'Prizes'!D{p})", align=S.CENTER)
    hr = 7 + C.N_PR + 1
    header_row(ws, hr, ["#", "Classroom", "Student", "Raised", "Highest prize earned", "Picked up?"], 1, height=22)
    first = hr + 1
    for k in range(1, C.N_STU + 1):
        r = hr + k
        put(ws, f"H{r}", f"=IFERROR(MATCH(SMALL({cs('AC')},{k}),{cs('AC')},0),\"\")", font=S.font(8, color="FFFFFF"))
        put(ws, f"A{r}", f"=IF(H{r}=\"\",\"\",{k})", align=S.CENTER)
        put(ws, f"B{r}", f"=IF(H{r}=\"\",\"\",INDEX({cs('G')},H{r}))")
        put(ws, f"C{r}", f"=IF(H{r}=\"\",\"\",INDEX({cs('AG')},H{r}))", font=S.font(11, True))
        put(ws, f"D{r}", f"=IF(H{r}=\"\",\"\",INDEX({cs('N')},H{r}))", fmt=S.FMT_USD2)
        put(ws, f"E{r}", f"=IF(H{r}=\"\",\"\",INDEX({cs('U')},H{r}))")
        put(ws, f"F{r}", f"=IF(H{r}=\"\",\"\",\"☐\")", align=S.CENTER)
    row_lines(ws, f"A{first}:F{hr + C.N_STU}", f"$H{first}")
    ws.column_dimensions["H"].hidden = True
    print_setup(ws, title_rows=f"{hr}:{hr}")
    dynamic_print_area(ws, "F", f"COUNT($H${first}:$H${hr + C.N_STU})", hr)
    ws.freeze_panes = f"A{first}"
    protect(ws)


def build_data_check(ws):
    ws.sheet_properties.tabColor = S.TAB_OUTPUT
    ws.sheet_view.showGridLines = False
    banner(ws, "A1:F1", "DATA CHECK  •  review after every update", 18, fill=S.NAVY)
    widths(ws, {"A": 4, "B": 46, "C": 16, "D": 16, "E": 40, "F": 14, "G": 6})
    put(ws, "B3", "=IF(C4=0,\"ALL CLEAR ✔\",\"CHECK \"&C4&\" ISSUE(S) BELOW\")", font=S.display(16, S.GREEN))
    ws.conditional_formatting.add("B3", FormulaRule(formula=["$C$4>0"], font=S.font(16, True, S.RED)))
    put(ws, "B4", "Issues found (rows flagged ✖ below)", font=S.font(10, italic=True))
    header_row(ws, 6, ["", "Check", "Count", "Amount", "What to do", "Status"], 1, height=20)
    o = lambda status: f"COUNTIF({oc('L')},\"{status}\")"  # noqa: E731
    osum = lambda status: f"SUMIFS({oc('H')},{oc('L')},\"{status}\")"  # noqa: E731
    checks = [
        ("OneCause rows pasted", f"=COUNTIF({oc('L')},\"?*\")", f"=SUM({oc('H')})", "", None),
        ("  Counted (matched, valid)", f"={o('OK')}", f"={osum('OK')}", "", None),
        ("  Unmatched student names", f"={o('UNMATCHED')}", f"={osum('UNMATCHED')}",
         "Add the OneCause name to Roster col H (OneCause name)", "C"),
        ("  Duplicate transaction IDs (ignored)", f"={o('DUPLICATE TXN')}", f"={osum('DUPLICATE TXN')}",
         "Usually a double paste; OK to ignore", None),
        ("  Negative / refunds (ignored)", f"={o('NEGATIVE / REFUND')}", f"={osum('NEGATIVE / REFUND')}",
         "Remove the refunded gift row from the paste", "C"),
        ("  Zero amounts (ignored)", f"={o('ZERO AMOUNT')}", "=0", "", None),
        ("  Bad dates (ignored)", f"={o('BAD DATE')}", f"={osum('BAD DATE')}", "Fix the date in the paste", "C"),
        ("  Bad amounts (ignored)", f"={o('BAD AMOUNT')}", "", "Fix the amount in the paste", "C"),
        ("OneCause reconciliation difference", "", "=ROUND(D7-SUM(D8:D13),2)",
         "Must be $0.00", "D"),
        ("Cash/check rows entered", f"=COUNTIF({cc('J')},\"?*\")", f"=SUM({cc('D')})", "", None),
        ("  Cash/check rows with problems", f"=COUNTIF({cc('J')},\"?*\")-COUNTIF({cc('J')},\"OK\")", "",
         "See list below; fix on Cash_Check", "C"),
        ("Gifts dated outside campaign (still counted)", f"=COUNTIFS({lg('F')},\"OK\",{lg('B')},\"<\"&StartDate)"
         f"+COUNTIFS({lg('F')},\"OK\",{lg('B')},\">\"&EndDate)", "='Calc_Daily'!P2+'Calc_Daily'!P3",
         "Check the dates are right", None),
        ("Totals tie-out (students = counted gifts)", "", f"=ROUND(Raised-SUM({lg('E')}),2)",
         "Must be $0.00", "D"),
        ("Roster rows needing attention", "=COUNTIF('Roster'!$K$2:$K$" + str(STU_LAST) + ",\"?*\")-COUNTIF("
         "'Roster'!$K$2:$K$" + str(STU_LAST) + ",\"✔\")", "", "See Roster col K", "C"),
        ("Duplicate classrooms", f"=SUMPRODUCT(({ccl('B')}<>\"\")*(COUNTIF({ccl('B')},{ccl('B')})>1))", "",
         "Each Teacher + Grade must be listed once on Classrooms", "C"),
        ("Prize thresholds out of order", f"=SUMPRODUCT(('Prizes'!$B$3:$B${PR_LAST}<>\"\")*"
         f"('Prizes'!$B$3:$B${PR_LAST}<='Prizes'!$B$2:$B${PR_LAST - 1}))", "",
         "Sort Prizes smallest to largest", "C"),
        ("OneCause required headers missing", "=COUNTIF('Settings'!$D$19:$D$20,\"✖*\")+IF(AND(OC_ColName=0,"
         "OR(OC_ColFirst=0,OC_ColLast=0)),1,0)", "", "Fix header names on Settings", "C"),
    ]
    flag_cells = []
    for i, (label, cnt, amt, todo, flag) in enumerate(checks):
        r = 7 + i
        put(ws, f"B{r}", label, border=S.BOX, font=S.font(11, not label.startswith("  ")))
        put(ws, f"C{r}", cnt or None, border=S.BOX, align=S.CENTER)
        put(ws, f"D{r}", amt or None, border=S.BOX, fmt=S.FMT_USD2)
        put(ws, f"E{r}", todo, border=S.BOX, font=S.font(10, italic=True))
        if flag == "C":
            put(ws, f"F{r}", f"=IF(N(C{r})=0,\"✔\",\"✖\")", border=S.BOX, align=S.CENTER)
            flag_cells.append(f"IF(N(C{r})=0,0,1)")
        elif flag == "D":
            put(ws, f"F{r}", f"=IF(ABS(N(D{r}))<0.005,\"✔\",\"✖\")", border=S.BOX, align=S.CENTER)
            flag_cells.append(f"IF(ABS(N(D{r}))<0.005,0,1)")
        else:
            put(ws, f"F{r}", "", border=S.BOX)
    put(ws, "C4", "=" + "+".join(flag_cells), font=S.font(10, True))
    ws.conditional_formatting.add("F7:F30", FormulaRule(formula=['F7="✖"'], font=S.font(12, True, S.RED),
                                                        fill=S.fill("F8D7DA")))
    ws.conditional_formatting.add("F7:F30", FormulaRule(formula=['F7="✔"'], font=S.font(12, True, S.GREEN)))

    # Problem lists
    def problem_list(row0, title, cols, rng_, bad_value, fields, n=C.N_LIST):
        """Rows of rng_ that are non-blank and not equal to bad_value's 'good' marker."""
        put(ws, f"B{row0}", title, font=S.display(12, S.RED))
        header_row(ws, row0 + 1, cols, 2, height=20)
        for k in range(1, n + 1):
            r = row0 + 1 + k
            arr(ws, f"G{r}", (f"=IFERROR(SMALL(IF(({rng_}<>\"{bad_value}\")*({rng_}<>\"\"),ROW({rng_})-1,999999),"
                              f"{k}),999999)"))
            ws[f"G{r}"].font = S.font(8, color="FFFFFF")
            for j, src in enumerate(fields):
                put(ws, f"{CL(2 + j)}{r}", f"=IF(G{r}=999999,\"\",{src.format(i=f'G{r}')})")
        row_lines(ws, f"B{row0 + 2}:E{row0 + 1 + n}", f"$B{row0 + 2}")
        return row0 + n + 3

    ro = lambda c: rng("Roster", c, 2, STU_LAST)  # noqa: E731
    nxt = problem_list(27, "OneCause rows not counted (first 40)",
                       ["Paste row #", "Status", "Student name in export", "Amount"], oc("L"), "OK",
                       ["{i}+1", f"INDEX({oc('L')},{{i}})", f"INDEX({oc('C')},{{i}})", f"INDEX({oc('B')},{{i}})"])
    nxt = problem_list(nxt, "Cash_Check rows with problems (first 40)",
                       ["Row #", "Status", "Student entry", "Amount"], cc("J"), "OK",
                       ["{i}+1", f"INDEX({cc('J')},{{i}})", f"INDEX({cc('B')},{{i}})", f"INDEX({cc('D')},{{i}})"])
    problem_list(nxt, "Roster rows needing attention (first 40)",
                 ["Row #", "Issue", "Student", "Teacher – Grade"], ro("K"), "✔",
                 ["{i}+1", f"INDEX({ro('K')},{{i}})", f"INDEX({ro('B')},{{i}})&\" \"&INDEX({ro('C')},{{i}})",
                  f"INDEX({ro('J')},{{i}})"])
    ws.column_dimensions["G"].hidden = True
    print_setup(ws)
    protect(ws)


def build_start_here(ws):
    ws.sheet_properties.tabColor = S.GOLD
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 3, "B": 6, "C": 70, "D": 26})
    banner(ws, "B1:D1", "SAINTS RUN WILD  •  START HERE", 22, fill=S.RED)
    put(ws, "B2", "=SchoolName&\"  •  Goal \"&TEXT(Goal,\"$#,##0\")&\"  •  \"&TEXT(StartDate,\"m/d\")&\" – \"&"
        "TEXT(EndDate,\"m/d/yyyy\")", font=S.font(12, True, S.NAVY))
    put(ws, "B3", "=\"Raised so far: \"&TEXT(Raised,\"$#,##0\")&\"  (\"&TEXT(PctGoal,\"0%\")&\")   •   Data Check: \"&"
        "'Data Check'!B3", font=S.font(12, True, S.GREEN))
    sections = [
        ("ONE-TIME SETUP (before kickoff)", [
            ("Settings", "Confirm goal, stretch milestone, dates, privacy mode, and the OneCause header names."),
            ("Classrooms", "List every Teacher + Grade once (Classroom name builds itself)."),
            ("Roster", "Paste students: First, Last, Grade, Teacher, Family. Student IDs are pre-assigned."),
            ("Prizes", "Enter milestone prizes smallest to largest."),
            ("Challenges", "Pre-filled from the Daily Incentives flyer. Adjust dates/rewards if needed."),
            ("History", "Optional: type last year's totals for the year-over-year comparison."),
        ]),
        ("EVERY DAY (about 5 minutes)", [
            ("OneCause_Paste", "1. Export the OneCause donation report (CSV/Excel). Select the whole sheet here, "
                               "Delete, then paste the export at A1 (headers in row 1)."),
            ("Cash_Check", "2. Type each cash/check gift: date, pick the student, amount, method."),
            ("Data Check", "3. Must say ALL CLEAR. Fix unmatched names by adding a OneCause name on the Roster."),
            ("Challenge Winners", "4. Read today's result; use Challenge Detail for per-grade/per-class winners."),
            ("TV Dashboard", "5. Save. Put this tab on the school TV (hide ribbon: Ctrl+F1, zoom to fit)."),
        ]),
        ("REPORTS (File → Export → PDF, or Print)", [
            ("Rpt_Summary", "One-page schoolwide summary"),
            ("Rpt_Grades", "Grade performance vs last year"),
            ("Rpt_Classes", "Every classroom: participation, totals, goal %"),
            ("Rpt_Students", "Every student total (admin only: full names)"),
            ("Rpt_Families", "Household totals"),
            ("Rpt_Donors", "Donor contributions (admin only)"),
            ("Rpt_Daily", "Daily and cumulative giving"),
            ("Prize Report", "Who earned which prize; pickup checklist"),
            ("Leaderboards", "Top students, classes, grades, families"),
        ]),
        ("AFTER THE FUN RUN", [
            ("History", "Copy this year's totals (Rpt_Summary + Rpt_Grades + Calc daily) into a new History row."),
            ("Settings", "Save As next year's file, clear Roster/OneCause_Paste/Cash_Check, update dates & goal."),
        ]),
    ]
    r = 5
    for title, items in sections:
        ws.merge_cells(f"B{r}:D{r}")
        put(ws, f"B{r}", title, font=S.display(13, S.WHITE), fill=S.NAVY)
        r += 1
        for tab, text in items:
            c = put(ws, f"D{r}", f"Open {tab} →", font=S.font(11, True, "1155CC"))
            c.hyperlink = f"#'{tab}'!A1"
            put(ws, f"C{r}", text, align=S.LEFT_WRAP)
            put(ws, f"B{r}", "•", align=S.CENTER, font=S.font(14, True, S.RED))
            ws.row_dimensions[r].height = 30
            r += 1
        r += 1
    put(ws, f"B{r}", "Color key:", font=S.font(11, True))
    put(ws, f"C{r}", "Sand cells = type here.  Gray-blue cells = automatic, don't type.  Blue tabs = inputs, red "
        "tabs = dashboards, green tabs = printable reports.", align=S.LEFT_WRAP)
    ws.row_dimensions[r].height = 32
    put(ws, f"C{r - 1}", fill=None)
    put(ws, f"D{r}", fill=S.INPUT)
    print_setup(ws)
    protect(ws)


# ----------------------------------------------------------------------------
# REPORT SHEETS
# ----------------------------------------------------------------------------
def report_title(ws, title, last_col):
    banner(ws, f"A1:{last_col}1", title, 16, fill=S.NAVY)
    ws.merge_cells(f"A2:{last_col}2")
    put(ws, "A2", "=CampaignName&\" \"&CampaignYear&\"  •  \"&SchoolName&\"  •  as of \"&TEXT(AsOfDate,\"m/d/yyyy\")",
        font=S.font(10, italic=True), align=S.CENTER)
    ws.sheet_properties.tabColor = S.TAB_REPORT
    ws.sheet_view.showGridLines = False


def build_rpt_summary(ws):
    report_title(ws, "FUNDRAISING SUMMARY", "H")
    widths(ws, {"A": 26, "B": 14, "C": 4, "D": 26, "E": 12, "F": 12, "G": 12, "H": 12})
    kpis = [("Raised", "=Raised", S.FMT_USD2), ("Goal", "=Goal", S.FMT_USD), ("Remaining", "=Remaining", S.FMT_USD2),
            ("% of goal", "=PctGoal", S.FMT_PCT1), ("Online (OneCause)", "=RaisedOnline", S.FMT_USD2),
            ("Cash / check", "=RaisedOffline", S.FMT_USD2), ("Students", "=Students", "0"),
            ("Participating students", "=Participants", "0"), ("Participation", "=PartPct", S.FMT_PCT1),
            ("Gifts", "=Gifts", "#,##0"), ("Unique donors", "=Donors", "#,##0"),
            ("Avg per student", "=AvgStudent", S.FMT_USD2), ("Last year total", "=LYTotal", S.FMT_USD),
            ("Data Check", "=IF(IssueCount=0,\"All clear\",IssueCount&\" issue(s) to review\")", None)]
    header_row(ws, 4, ["Schoolwide", "Value"], 1, height=20)
    for i, (lab, f, fmt) in enumerate(kpis):
        put(ws, f"A{5 + i}", lab, border=S.BOX, font=S.font(11, True))
        put(ws, f"B{5 + i}", f, border=S.BOX, fmt=fmt, align=S.RIGHT)
    header_row(ws, 4, ["Grade", "Raised", "Particip.", "Students", "Last year"], 4, height=20)
    for i in range(len(C.GRADES)):
        r = 5 + i
        g = i + 2
        put(ws, f"D{r}", f"=\"Grade \"&'Calc_Grades'!A{g}", border=S.BOX)
        put(ws, f"E{r}", f"='Calc_Grades'!G{g}", border=S.BOX, fmt=S.FMT_USD)
        put(ws, f"F{r}", f"='Calc_Grades'!F{g}", border=S.BOX, fmt=S.FMT_PCT)
        put(ws, f"G{r}", f"='Calc_Grades'!D{g}", border=S.BOX, fmt="0")
        put(ws, f"H{r}", f"='Calc_Grades'!K{g}", border=S.BOX, fmt=S.FMT_USD)
    r0 = 21
    header_row(ws, r0, ["Top classrooms", "Raised"], 1, height=20)
    header_row(ws, r0, ["Top students (full name)", "Classroom", "", "", "Raised"], 4, height=20)
    for k in range(1, 11):
        r = r0 + k
        v = f"LARGE({ccl('M')},{k})"
        put(ws, f"A{r}", f"=IF({v}<0.005,\"\",INDEX({ccl('B')},MATCH({v},{ccl('M')},0)))", border=S.BOX)
        put(ws, f"B{r}", f"=IF({v}<0.005,\"\",INDEX({ccl('I')},MATCH({v},{ccl('M')},0)))", border=S.BOX,
            fmt=S.FMT_USD)
        sv = f"LARGE({cs('Q')},{k})"
        put(ws, f"D{r}", f"=IF({sv}<0.005,\"\",INDEX({cs('AG')},MATCH({sv},{cs('Q')},0)))", border=S.BOX)
        ws.merge_cells(f"E{r}:G{r}")
        put(ws, f"E{r}", f"=IF({sv}<0.005,\"\",INDEX({cs('G')},MATCH({sv},{cs('Q')},0)))", border=S.BOX)
        put(ws, f"H{r}", f"=IF({sv}<0.005,\"\",INDEX({cs('N')},MATCH({sv},{cs('Q')},0)))", border=S.BOX,
            fmt=S.FMT_USD)
    r0 = 33
    header_row(ws, r0, ["Challenge", "Status", "", "Result"], 1, height=20)
    for i in range(C.N_CH):
        r = r0 + 1 + i
        k = i + 2
        put(ws, f"A{r}", f"='Calc_Challenges'!D{k}", border=S.BOX)
        put(ws, f"B{r}", f"='Calc_Challenges'!K{k}", border=S.BOX)
        ws.merge_cells(f"D{r}:H{r}")
        put(ws, f"D{r}", f"='Calc_Challenges'!Y{k}", border=S.BOX)
    r0 = 50
    header_row(ws, r0, ["Prize", "Earned"], 1, height=20)
    for i in range(C.N_PR):
        put(ws, f"A{r0 + 1 + i}", f"=IF('Prizes'!A{i + 2}=\"\",\"\",'Prizes'!A{i + 2})", border=S.BOX)
        put(ws, f"B{r0 + 1 + i}", f"='Prizes'!D{i + 2}", border=S.BOX, align=S.CENTER)
    print_setup(ws)
    ws.page_setup.fitToHeight = 1
    ws.print_area = "A1:H62"


def build_rpt_grades(ws):
    report_title(ws, "GRADE PERFORMANCE", "J")
    widths(ws, {"A": 12, "B": 9, "C": 10, "D": 13, "E": 13, "F": 14, "G": 14, "H": 14, "I": 14, "J": 9})
    header_row(ws, 4, ["Grade", "Classes", "Students", "Participants", "Participation", "Raised", "Avg / student",
                       "Last year", "Change", "Rank"], 1)
    for i in range(len(C.GRADES)):
        r, g = 5 + i, i + 2
        vals = [(f"=\"Grade \"&'Calc_Grades'!A{g}", None), (f"='Calc_Grades'!C{g}", "0"),
                (f"='Calc_Grades'!D{g}", "0"), (f"='Calc_Grades'!E{g}", "0"), (f"='Calc_Grades'!F{g}", S.FMT_PCT1),
                (f"='Calc_Grades'!G{g}", S.FMT_USD2), (f"='Calc_Grades'!H{g}", S.FMT_USD2),
                (f"='Calc_Grades'!K{g}", S.FMT_USD), (f"='Calc_Grades'!L{g}", '+"$"#,##0;-"$"#,##0;"$"0'),
                (f"='Calc_Grades'!J{g}", "0")]
        for j, (f, fmt) in enumerate(vals):
            put(ws, f"{CL(1 + j)}{r}", f, fmt=fmt, border=S.BOX)
    r = 5 + len(C.GRADES)
    put(ws, f"A{r}", "TOTAL", font=S.font(11, True), border=S.BOX)
    for col, f, fmt in [("B", "=SUM(B5:B15)", "0"), ("C", "=SUM(C5:C15)", "0"), ("D", "=SUM(D5:D15)", "0"),
                        ("E", "=IF(C16>0,D16/C16,0)", S.FMT_PCT1), ("F", "=SUM(F5:F15)", S.FMT_USD2),
                        ("G", "=IF(C16>0,F16/C16,0)", S.FMT_USD2), ("H", "=SUM(H5:H15)", S.FMT_USD),
                        ("I", "=F16-H16", '+"$"#,##0;-"$"#,##0;"$"0')]:
        put(ws, f"{col}{r}", f, fmt=fmt, border=S.BOX, font=S.font(11, True))
    bar = BarChart()
    bar.type = "bar"
    bar.title = "Raised by grade"
    bar.add_data(Reference(ws.parent["Calc_Grades"], min_col=7, min_row=1, max_row=G_LAST), titles_from_data=True)
    bar.set_categories(Reference(ws.parent["Calc_Grades"], min_col=13, min_row=2, max_row=G_LAST))
    bar.series[0].graphicalProperties.solidFill = S.RED
    bar.legend = None
    bar.width, bar.height = 18, 9
    bar.x_axis.scaling.orientation = "maxMin"
    bar.y_axis.numFmt = '"$"#,##0'
    bar.visible_cells_only = False
    ws.add_chart(bar, "A19")
    print_setup(ws)
    ws.print_area = "A1:J38"


def build_long_report(ws, title, cols, widths_spec, n, key_formula, fields, last_col, landscape=False,
                      total_row=None):
    """Sorted list report: helper col Z holds the source row index of the k-th item."""
    report_title(ws, title, last_col)
    widths(ws, widths_spec)
    header_row(ws, 4, cols, 1)
    for k in range(1, n + 1):
        r = 4 + k
        put(ws, f"Z{r}", key_formula.format(k=k), font=S.font(8, color="FFFFFF"))
        for j, (src, fmt) in enumerate(fields):
            col = CL(1 + j)
            put(ws, f"{col}{r}", f"=IF(Z{r}=\"\",\"\",{src.format(i=f'Z{r}', k=k)})", fmt=fmt)
    row_lines(ws, f"A5:{last_col}{4 + n}", "$Z5")
    ws.column_dimensions["Z"].hidden = True
    ws.freeze_panes = "A5"
    print_setup(ws, landscape=landscape, title_rows="4:4")
    dynamic_print_area(ws, last_col, f"COUNT($Z$5:$Z${4 + n})", 4)


def build_rpt_classes(ws):
    build_long_report(
        ws, "CLASSROOM FUNDRAISING & PARTICIPATION",
        ["Grade", "Classroom", "Students", "Participants", "Participation", "Raised", "Avg / student", "Class goal",
         "% of goal", "Rank $"],
        {"A": 8, "B": 30, "C": 10, "D": 12, "E": 13, "F": 13, "G": 13, "H": 12, "I": 10, "J": 8},
        C.N_CLS, f"=IFERROR(MATCH(SMALL({ccl('S')},{{k}}),{ccl('S')},0),\"\")",
        [(f"INDEX({ccl('D')},{{i}})", None), (f"INDEX({ccl('B')},{{i}})", None), (f"INDEX({ccl('F')},{{i}})", "0"),
         (f"INDEX({ccl('G')},{{i}})", "0"), (f"INDEX({ccl('H')},{{i}})", S.FMT_PCT1),
         (f"INDEX({ccl('I')},{{i}})", S.FMT_USD2), (f"INDEX({ccl('J')},{{i}})", S.FMT_USD2),
         (f"IF(INDEX({ccl('K')},{{i}})=0,\"\",INDEX({ccl('K')},{{i}}))", S.FMT_USD),
         (f"INDEX({ccl('L')},{{i}})", S.FMT_PCT), (f"INDEX({ccl('O')},{{i}})", "0")],
        "J")


def build_rpt_students(ws):
    build_long_report(
        ws, "STUDENT FUNDRAISING TOTALS (CONFIDENTIAL)",
        ["Classroom", "Student", "Online", "Cash/check", "Total", "Gifts", "Highest prize", "$ to next prize"],
        {"A": 28, "B": 26, "C": 11, "D": 11, "E": 11, "F": 7, "G": 22, "H": 13},
        C.N_STU, f"=IFERROR(MATCH(SMALL({cs('AD')},{{k}}),{cs('AD')},0),\"\")",
        [(f"INDEX({cs('G')},{{i}})", None), (f"INDEX({cs('D')},{{i}})&\", \"&INDEX({cs('C')},{{i}})", None),
         (f"INDEX({cs('L')},{{i}})", S.FMT_USD2), (f"INDEX({cs('M')},{{i}})", S.FMT_USD2),
         (f"INDEX({cs('N')},{{i}})", S.FMT_USD2), (f"INDEX({cs('O')},{{i}})", "0"),
         (f"INDEX({cs('U')},{{i}})", None), (f"INDEX({cs('W')},{{i}})", S.FMT_USD2)],
        "H", landscape=False)


def build_rpt_families(ws):
    fr = lambda c: rng("Calc_Families", c, 2, FAM_LAST)  # noqa: E731
    build_long_report(
        ws, "FAMILY / HOUSEHOLD TOTALS (CONFIDENTIAL)",
        ["Rank", "Family", "Students", "Participating", "Raised"],
        {"A": 7, "B": 32, "C": 10, "D": 13, "E": 13},
        C.N_FAM, f"=IFERROR(IF(LARGE({fr('F')},{{k}})<0,\"\",MATCH(LARGE({fr('F')},{{k}}),{fr('F')},0)),\"\")",
        [(f"INDEX({fr('G')},{{i}})", "0"), (f"INDEX({fr('B')},{{i}})", None), (f"INDEX({fr('C')},{{i}})", "0"),
         (f"INDEX({fr('D')},{{i}})", "0"), (f"INDEX({fr('E')},{{i}})", S.FMT_USD2)],
        "E")


def build_rpt_donors(ws):
    dr = lambda c: rng("Calc_Donors", c, 2, DON_LAST)  # noqa: E731
    build_long_report(
        ws, "DONOR CONTRIBUTIONS (CONFIDENTIAL)",
        ["#", "Donor", "Gifts", "Total given"],
        {"A": 6, "B": 36, "C": 8, "D": 14},
        C.N_DON, f"=IFERROR(IF(LARGE({dr('F')},{{k}})<0,\"\",MATCH(LARGE({dr('F')},{{k}}),{dr('F')},0)),\"\")",
        [("{k}", "0"), (f"INDEX({dr('B')},{{i}})", None), (f"INDEX({dr('D')},{{i}})", "0"),
         (f"INDEX({dr('E')},{{i}})", S.FMT_USD2)],
        "D")
    put(ws, "F4", "Donor names come from OneCause 'Donor name' and Cash_Check 'Donor / family'.",
        font=S.font(9, italic=True))


def build_rpt_daily(ws, days):
    report_title(ws, "DAILY GIVING", "H")
    widths(ws, {"A": 8, "B": 14, "C": 12, "D": 12, "E": 12, "F": 8, "G": 14, "H": 14})
    header_row(ws, 4, ["Day #", "Date", "Online", "Cash/check", "Total", "Gifts", "Cumulative", "Participants"], 1)
    for d in range(1, C.N_DAY + 1):
        r, s = 4 + d, 1 + d
        for col, src, fmt in [("A", "A", "0"), ("B", "B", "ddd m/d"), ("C", "D", S.FMT_USD2),
                              ("D", "E", S.FMT_USD2), ("E", "F", S.FMT_USD2), ("F", "G", "0"),
                              ("G", "H", S.FMT_USD2), ("H", "I", "0")]:
            put(ws, f"{col}{r}", f"=IF('Calc_Daily'!$B{s}=\"\",\"\",'Calc_Daily'!{src}{s})", fmt=fmt)
    row_lines(ws, f"A5:H{4 + C.N_DAY}", "$B5")
    r = 5 + C.N_DAY
    put(ws, f"B{r}", "Before start", font=S.font(10, italic=True))
    put(ws, f"E{r}", "='Calc_Daily'!P2", fmt=S.FMT_USD2)
    put(ws, f"B{r + 1}", "After end", font=S.font(10, italic=True))
    put(ws, f"E{r + 1}", "='Calc_Daily'!P3", fmt=S.FMT_USD2)
    dws = ws.parent["Calc_Daily"]
    bar = BarChart()
    bar.title = "Raised per day"
    bar.add_data(Reference(dws, min_col=6, min_row=1, max_row=1 + days), titles_from_data=True)
    bar.set_categories(Reference(dws, min_col=13, min_row=2, max_row=1 + days))
    bar.series[0].graphicalProperties.solidFill = S.NAVY
    bar.legend = None
    bar.y_axis.numFmt = '"$"#,##0'
    bar.width, bar.height = 18, 8
    bar.visible_cells_only = False
    ws.add_chart(bar, "J4")
    print_setup(ws, landscape=True)


# ----------------------------------------------------------------------------
# ASSEMBLY
# ----------------------------------------------------------------------------
ORDER = ["Start Here", "TV Dashboard", "Leaderboards", "Challenge Winners", "Prize Report", "Data Check",
         "Settings", "Classrooms", "Roster", "Prizes", "Challenges", "OneCause_Paste", "Cash_Check", "History",
         "Rpt_Summary", "Rpt_Grades", "Rpt_Classes", "Rpt_Students", "Rpt_Families", "Rpt_Donors", "Rpt_Daily",
         "OC_Norm", "Ledger", "Calc_Students", "Calc_Classes", "Calc_Grades", "Calc_Families", "Calc_Donors",
         "Calc_Daily", "Calc_Challenges", "Calc_KPI"]
ENGINE = ORDER[ORDER.index("OC_Norm"):]


def build(path, data=None):
    data = data or {}
    settings = dict(C.SETTINGS, **data.get("settings", {}))
    prizes = data.get("prizes", C.PRIZES)
    challenges = data.get("challenges", C.CHALLENGES)
    oc_headers = data.get("oc_headers", ["Transaction Date", "Transaction ID", "Participant Name", "Donor Name",
                                         "Amount", "Payment Method"])

    wb = Workbook()
    wb.remove(wb.active)
    sheets = {n: wb.create_sheet(n) for n in ORDER}
    name(wb, "LB_Grade", "'Leaderboards'!$C$3")

    build_settings(wb, sheets["Settings"], settings)
    build_classrooms(sheets["Classrooms"], data.get("classrooms", []))
    build_roster(sheets["Roster"], data.get("roster", []))
    build_prizes(sheets["Prizes"], prizes)
    build_challenges(sheets["Challenges"], challenges)
    build_onecause_paste(sheets["OneCause_Paste"], oc_headers, data.get("oc_rows", []))
    build_cash(sheets["Cash_Check"], data.get("cash", []))
    build_history(sheets["History"], data.get("history", {}))

    build_oc_norm(sheets["OC_Norm"])
    build_ledger(sheets["Ledger"])
    build_calc_students(sheets["Calc_Students"])
    build_calc_classes(sheets["Calc_Classes"])
    build_calc_grades(sheets["Calc_Grades"])
    build_calc_families(sheets["Calc_Families"])
    build_calc_donors(sheets["Calc_Donors"])
    build_calc_daily(sheets["Calc_Daily"])
    build_calc_challenges(sheets["Calc_Challenges"])
    build_calc_kpi(wb, sheets["Calc_KPI"])

    pick_i = data.get("challenge_pick", 1) - 1
    pc = challenges[pick_i] if challenges else None
    default_pick = f"{pick_i + 1}. {pc[3]} ({pc[0].month}/{pc[0].day})" if pc else ""
    build_start_here(sheets["Start Here"])
    days = min(C.N_DAY, (settings["end"] - settings["start"]).days + 1)
    build_tv(sheets["TV Dashboard"], days)
    build_leaderboards(sheets["Leaderboards"])
    build_challenge_winners(sheets["Challenge Winners"], default_pick)
    build_prize_report(sheets["Prize Report"], data.get("prize_pick", prizes[0][0] if prizes else ""))
    build_data_check(sheets["Data Check"])
    build_rpt_summary(sheets["Rpt_Summary"])
    build_rpt_grades(sheets["Rpt_Grades"])
    build_rpt_classes(sheets["Rpt_Classes"])
    build_rpt_students(sheets["Rpt_Students"])
    build_rpt_families(sheets["Rpt_Families"])
    build_rpt_donors(sheets["Rpt_Donors"])
    build_rpt_daily(sheets["Rpt_Daily"], days)

    for n in ENGINE:
        sheets[n].sheet_state = "hidden"
        protect(sheets[n])
    if data.get("preview_only"):  # render check: show just these sheets
        keep = data["preview_only"]
        wb.active = ORDER.index(keep[0])
        for n in ORDER:
            if n not in keep:
                sheets[n].sheet_state = "hidden"
    else:
        wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=os.path.join(os.path.dirname(__file__), "..", "output"))
    ap.add_argument("--blank", action="store_true", help="build only the blank template")
    args = ap.parse_args()
    year = C.SETTINGS["year"]
    print("built", build(os.path.join(args.out, f"SaintsRunWild_{year}.xlsx")))
    if not args.blank:
        import demo_data
        print("built", build(os.path.join(args.out, f"SaintsRunWild_{year}_DEMO.xlsx"), demo_data.generate()))


if __name__ == "__main__":
    main()
